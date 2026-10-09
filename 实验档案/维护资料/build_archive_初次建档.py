from pathlib import Path
import ast,csv,hashlib,json,re,shutil,subprocess,urllib.parse
from collections import deque
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(r'E:\DCT-Reg')
TASK=ROOT/'.review/experiment-ledger-20261009'
SRC=TASK/'source_snapshot'
DEST=ROOT/'实验档案'
DOCX=DEST/'DCT实验总档案.docx'
CODE=DEST/'实验代码'
ASSETS=DEST/'表格图片'
MAINT=DEST/'维护资料'
COMMIT='2b1c622f418c619dc20f973ddeb839ff12d30d88'
SNAP=CODE/'共享代码快照_2b1c622'
DATE='2026-10-09'
if DOCX.exists():raise SystemExit('Main Word already exists; do not overwrite manual additions.')
for p in [DEST,CODE,ASSETS,MAINT,SNAP]:p.mkdir(parents=True,exist_ok=True)

def load(rel):return json.loads((SRC/rel).read_text(encoding='utf8'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def clean(s):return s.replace('**','').strip()
md=(SRC/'paper/DCT_v313_初稿.md').read_text(encoding='utf8')
tables=[];block=[]
for line in md.splitlines()+['']:
    if line.startswith('|'):block.append([clean(x) for x in line.strip().strip('|').split('|')])
    elif block:
        tables.append([block[0]]+block[2:]);block=[]
main=next(t for t in tables if t[0]==['队列','Fold 0','Fold 1','Fold 2','Fold 3','Fold 4','均值 ± 标准差'])
weights=next(t for t in tables if t[0][0]=='配置' and 'IPCW 排序' in t[0])
abla=next(t for t in tables if t[0][0]=='配置' and 'BLCA 相对 Exp0' in t[0])
epoch_blca=next(t for t in tables if t[0][0]=='BLCA 配置')
epoch_kirc=next(t for t in tables if t[0][0]=='KIRC 配置')
vals=load('paper/figures/v313_manuscript_integrated/source_values.json')
audit=load('results/v313_evidence_v2/controls_audit.json')
assert audit['passed'] is True and audit['errors']==[] and len(audit['runs'])==20
meta=load('paper/figures/v313_manuscript_integrated/manifest.json')
profile=load('paper/figures/v313_tradeoff_real_20261008/tradeoff_blca_kirc/efficiency_tradeoff.json')
profile10=load('paper/figures/v313_tradeoff_real_20261008/tradeoff_ten_cohort_full_only/efficiency_per_cohort.json')
km={c:load(f'paper/figures/fig5_ten_cancers/km_{c.lower()}_exp6_oof.json') for c in [r[0] for r in main[1:]]}
case=load('paper/figures/v313_reference_panels_real_20261008/raw/pathway_case_panel_raw.json')
case_a=np.array(case['raw_attention'])
comparison=load('paper/V313_OT_UNI2H_COMPARISON_INPUT.json')
records=[];manifest=[]

def add_record(num,slug,title,kind,status,purpose,setting,conclusion,pending,sources,scripts):
    rec=dict(id=f'E{num:03d}',folder=f'E{num:03d}_{slug}',title=title,kind=kind,status=status,purpose=purpose,setting=setting,conclusion=conclusion,pending=pending,sources=sources,scripts=scripts,tables=[],figures=[],assets=[])
    records.append(rec)
    (ASSETS/rec['folder']).mkdir(parents=True,exist_ok=True)
    (CODE/rec['folder']).mkdir(parents=True,exist_ok=True)
    return rec

def table(rec,title,headers,rows):
    data=dict(title=title,headers=headers,rows=[[str(x) for x in row] for row in rows])
    rec['tables'].append(data)
    p=ASSETS/rec['folder']/f'表{len(rec["tables"]):02d}.csv'
    with p.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows([headers]+rows)
    return data

def file_asset(rec,rel,target=None,origin='commit'):
    p=SRC/rel if origin=='commit' else Path(rel)
    out=ASSETS/rec['folder']/'原始材料'/(target or p.name)
    out.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(p,out)
    ent=dict(experiment=rec['id'],source=rel,source_kind=origin,source_commit=COMMIT if origin=='commit' else 'local figure revision 2026-10-09 from 2b1c622',target=str(out.relative_to(DEST)),sha256=sha(out))
    assert sha(p)==ent['sha256']
    manifest.append(ent);rec['assets'].append(ent)
    return out

def folder_asset(rec,rel):
    root=SRC/rel
    assert root.exists(),rel
    result={}
    for p in sorted(root.rglob('*')):
        if p.is_file():result[p.name]=file_asset(rec,str(p.relative_to(SRC)).replace('\\','/'),str(Path(root.name)/p.relative_to(root)))
    return result

def figure(rec,path,caption):rec['figures'].append(dict(path=str(path),caption=caption))
LOSS_SOURCE='paper/figures/v313_manuscript_integrated/source_values.json'
LOSS_IMG='paper/figures/v313_manuscript_integrated/fig2_loss_ablation.png'
COMMON_SCRIPTS=['scripts/run_ablation_loss_components.py','survot_rank/training/train_runner.py','configs/dct_v313_uni2h.yaml']
MODEL='survot_rank/research/methods/legacy/experimental/dct_v313_transport_reconstruction/model.py'
loss_titles=['患者级 NLL 基线','增加 IPCW 排序','增加逐槽生存监督','增加槽多样性','仅加入自重建','仅加入运输跨模态重建','完整联合目标与十队列主结果']
loss_conclusions=[
'这是同架构训练目标基线。BLCA 0.7001、KIRC 0.8120，后续损失配置均应分别对照各癌种。',
'相对 Exp0，BLCA 均值增加 0.0096，KIRC 降低 0.0185。排序项没有在这两个队列上都带来提高。',
'相对 Exp1，BLCA 均值继续提高，KIRC 略降。逐槽监督的价值需结合槽退化诊断，不能只看槽数量。',
'相对 Exp2，两个队列均值均提高；相对 Exp0，KIRC 仍低于基线。这是当前单 seed 配方的观察。',
'以 Exp3 为起点，仅加入 self 重建，BLCA/KIRC 均值均未超过 Exp3。当前 self 有效系数为 0.025，不等于 Full 的 0.05。',
'以 Exp3 为起点，仅加入 cross 重建，均值未超过 Full。Exp4 与 Exp5 是并列支路，不能当作连续累加；当前 cross 系数为 0.025。',
'Full 在 BLCA/KIRC 为七组配置中的最高均值，相对 Exp0 分别提高 0.0237 和 0.0104。十队列等权平均约 0.7030，均为最佳验证表现。'
]
for i,title in enumerate(loss_titles):
    w=weights[i+1][1:]
    setting=f'共同五折开发协议，seed=3；有效系数为 IPCW {w[0]}，逐槽 NLL {w[1]}，多样性 {w[2]}，self {w[3]}，cross {w[4]}。重建系数为 ramp 达到上限后的有效值。'
    pending='补齐历史原始配置与患者预测的对应；不得用当前脚本默认值替代历史权重。Exp4/Exp5 与 Full 的单支权重不同，尚不能据此证明严格协同效应。'
    rec=add_record(i+1,f'Exp{i}',f'Exp{i} {title}','训练目标消融' if i<6 else '完整模型训练与跨队列汇总','已有五折记录',f'检验{title}在同架构中的结果。',setting,loss_conclusions[i],pending,[LOSS_SOURCE,'paper/DCT_v313_初稿.md §4.4 与附录 A'],COMMON_SCRIPTS+(['scripts/run_v313_uni2h_10cancer.py','scripts/summarize_v313_exp6.py'] if i==6 else []))
    rows=[]
    for c in ['blca','kirc']:
        x=np.array(vals['loss_folds'][c][i]);base=np.array(vals['loss_folds'][c][0])
        rows.append([c.upper(),f'{x.mean():.4f} ± {x.std(ddof=1):.4f}',f'{x.mean()-base.mean():+.4f}' if i else '基线'])
    table(rec,'当前消融均值与基线差',['队列','五折均值 ± 样本标准差','相对 Exp0'],rows)
    table(rec,'逐折分数与最佳 epoch',['队列','Fold 0','Fold 1','Fold 2','Fold 3','Fold 4'],[['BLCA']+epoch_blca[i+1][1:],['KIRC']+epoch_kirc[i+1][1:]])
    rec['note']='括号内为最佳 epoch；本表保留历史四位记录。均值和样本标准差按对应四位折值计算。'
    file_asset(rec,LOSS_SOURCE)
    loss_image=file_asset(rec,LOSS_IMG)
    if i==6:
        rows=[[r[0],r[-1],str(km[r[0]]['n'])] for r in main[1:]]
        table(rec,'Full 十队列主结果',['队列','五折 C-index 均值 ± 标准差','验证患者数'],rows)
        table(rec,'十队列完整逐折记录',main[0],main[1:])
        figure(rec,loss_image,'七组训练目标的联合对比图。Full 与各目标配置来自同一消融记录；图不是七次独立验证。')
        folder_asset(rec,'paper/tables/v313_training_variants_20261008')

for num,arm,title in [(8,'direct','Direct 重建输入对照'),(9,'independent','Independent 学习耦合对照')]:
    rec=add_record(num,arm,title,'重训机制对照','服务器审计记录通过',
      '检验辅助重建读取运输对齐信息的作用。' if arm=='direct' else '检验非独立联合耦合在预测和重建路径中的作用。',
      'BLCA/KIRC 各五折，seed=3，legacy_val；self/cross 有效系数各为 0.05。Direct 保留预测 OT，仅将 cross 输入改为原始病理槽。' if arm=='direct' else 'BLCA/KIRC 各五折；用学习边际的乘积替换联合耦合，并同时用于预测和 cross 重建。',
      'Full 相对 Direct 在 BLCA/KIRC 分别提高约 0.0209/0.0193；逐折提高比例为 3/5 与 4/5。' if arm=='direct' else 'Full 相对 Independent 在 BLCA/KIRC 分别提高约 0.0146/0.0170；逐折提高比例为 3/5 与 5/5。',
      '审计来自已提交服务器记录，本次整理不重跑模型。Full 与控制臂的完整患者、结局、配置配对核对仍应保留。Independent 同时改变两个位置，不能把差值只归因于重建。',
      ['results/v313_evidence_v2/controls.json','results/v313_evidence_v2/controls_audit.json'],
      ['scripts/prepare_v313_evidence.py','scripts/run_v313_additional_evidence.py','survot_rank/evidence/v313.py','survot_rank/evidence/additional.py',f'results/v313_evidence_v2/controls_plan_configs/{arm}_blca.yaml',f'results/v313_evidence_v2/controls_plan_configs/{arm}_kirc.yaml'])
    rows=[];fold_rows=[]
    for c in ['blca','kirc']:
        runs=sorted([v for v in audit['runs'] if v['arm']==arm and v['cancer']==c],key=lambda v:v['fold'])
        assert len(runs)==5 and all(not v['errors'] for v in runs)
        x=np.array([v['cindex'] for v in runs]);full=np.array(vals['full_exact_folds'][c])
        rows.append([c.upper(),f'{x.mean():.4f} ± {x.std(ddof=1):.4f}',f'{full.mean()-x.mean():+.4f}',f'{int((full>x).sum())}/5'])
        fold_rows.append([c.upper()]+[f"{v['cindex']:.4f} ({v['best_epoch']})" for v in runs])
    table(rec,'重训控制臂与 Full 对照',['队列','控制臂 C-index 均值 ± 标准差','Full 减控制臂','Full 较高折数'],rows)
    table(rec,'控制臂逐折分数与最佳 epoch',['队列','Fold 0','Fold 1','Fold 2','Fold 3','Fold 4'],fold_rows)
    for f in rec['sources']:file_asset(rec,f)
    for c in ['blca','kirc']:file_asset(rec,f'results/v313_evidence_v2/controls_plan_configs/{arm}_{c}.yaml')
    if num==9:
        f=file_asset(rec,'paper/figures/v313_manuscript_integrated/fig3_trained_controls.png')
        figure(rec,f,'Full 与两种重训控制臂的逐折对比；均值提高不代表每一折均提高。')

rec=add_record(10,'运输计划干预','固定 Full 模型的运输计划替换','固定 checkpoint 分析','十折 JSON 已保存','固定训练完成的 Full，检验同实际边缘的计划替换是否影响排序评价。','BLCA/KIRC 各五折，alpha 为 0、0.25、0.5、0.75、1；全部阶段与几何分支替换计划，重新计算后续门控。','BLCA 在全部强度下 C-index 不变；KIRC 仅 fold2 在非零强度下小幅变化。该固定模型的排序评价对这类替换基本不敏感。','尚缺训练时最佳患者预测的逐患者对齐、风险差分、重建误差和实际边缘残差。不得用旧三折汇总替代此十折记录。',['paper/figures/fig3_sweep_blca_fold0.json 至 fold4','paper/figures/fig3_sweep_kirc_fold0.json 至 fold4'],['scripts/plot_v313_manuscript_sweep.py','scripts/plot_fig3_from_exports.py','scripts/plot_fig3_transport_sweep.py','scripts/prepare_v313_evidence.py','survot_rank/evidence/v313.py'])
sweeps={}
for c in ['blca','kirc']:
    sweeps[c]=[load(f'paper/figures/fig3_sweep_{c}_fold{f}.json') for f in range(5)]
    for f in range(5):file_asset(rec,f'paper/figures/fig3_sweep_{c}_fold{f}.json')
rows=[]
for ai,alpha in enumerate(sweeps['blca'][0]['alphas']):
    vals_={c:np.array([v['cindex'][ai] for v in sweeps[c]]) for c in sweeps}
    rows.append([str(alpha)]+[f'{vals_[c].mean():.6f} ± {vals_[c].std(ddof=1):.6f}' for c in ['blca','kirc']])
table(rec,'计划替换的五折汇总',['alpha','BLCA C-index','KIRC C-index'],rows)
changed=sum(abs(v['cindex'][ai]-v['cindex'][0])>1e-12 for c in sweeps for v in sweeps[c] for ai in range(1,5))
assert changed==4
f=file_asset(rec,'paper/figures/fig4_transport_sweep_manuscript.png');figure(rec,f,'同一 Full 模型内的计划替换；不同于 E009 重新训练 Independent。')

rec=add_record(11,'十队列KM','十队列高低风险 KM 分层','已有预测的生存曲线分析','十队列材料齐全','检验 Full 的验证预测能否呈现一致方向的高低风险分层。','使用每折训练患者风险中位数对该折验证患者分组，再合并五折；包含删失标记、在险人数和名义 95% 区间。','10 个队列共 4760 名验证患者，均呈现高风险组总体生存曲线较低的方向；探索性 log-rank p 值均低于 0.05。','p 值和区间未校正交叉验证模型依赖；属于开发预测分层，不是独立外层验证。总览第二页存在排版问题，UCEC 图例与曲线重叠待修。',['paper/figures/fig5_ten_cancers/figure5_ten_cancers.json'],['scripts/plot_fig5_km_curves.py','scripts/rerun_v313_export_with_km.py','survot_rank/evidence/km_oof.py'])
assert sum(v['n'] for v in km.values())==4760
km_table=next(t for t in tables if t[0][0]=='队列' and '探索性 log-rank p' in t[0])
for row in km_table[1:]:
    j=km[row[0]];assert [int(v) for v in row[1:4]]==[j['n'],j['low_n'],j['high_n']]
table(rec,'十队列 KM 人数与探索性统计',km_table[0],km_table[1:])
files=folder_asset(rec,'paper/figures/fig5_ten_cancers')
figure(rec,files['km_blca_exp6_oof.png'],'BLCA 单图用于展示风险曲线、删失与在险人数；其他九队列单图及两张总览在材料目录。')

rec=add_record(12,'病例通路','BLCA 病例槽与通路解释','已有模型的病例分析','已有导出和图表 校版中','记录病例 TCGA-2F-A9KP 的八个组学槽与 329 条通路的最终聚合权重。','BLCA fold1、Exp6 Full、seed=3；展示 raw 权重、槽内百分位与 Top-3/Bottom-3；原始组织图和 Top-5 组织块未接入。','最终通路权重约 0.00237–0.00342，每槽权重和约 1，接近均匀分配 1/329。可描述相对偏好，当前图不能单独证明强生物学分工。','旧 raw 图从零起色标导致整体偏红；percentile 只表示排序。Top/Bottom 百分位条形图信息有限。尚缺真实组织空间证据；单病例聚合权重不能作为槽塌缩判据。',['paper/figures/v313_reference_panels_real_20261008/raw/pathway_case_panel_raw.json'],['scripts/plot_v313_slotspe_style.py','survot_rank/evidence/slotspe_style.py','scripts/rerun_v313_export_for_attention.py'])
prob=case_a/case_a.sum(1,keepdims=True);entropy=-(prob*np.log(prob)).sum(1)/np.log(329)
table(rec,'每槽原始权重诊断',['槽','最小权重','最大权重','标准化熵','Top-3 权重和'],[[str(s),f'{case_a[s].min():.6f}',f'{case_a[s].max():.6f}',f'{entropy[s]:.6f}',f'{np.sort(case_a[s])[-3:].sum():.6f}'] for s in range(8)])
folder_asset(rec,'paper/figures/v313_reference_panels_real_20261008')
revision=Path(r'C:\Users\栋栋\.codex\visualizations\2026\10\08\01a11aaf-3037-70b3-9649-30f9bfa9cf94\dct-paper-2b1c622\redesign_20261009')
for p in revision.iterdir():
    if p.suffix in ['.png','.pdf','.svg','.json','.py']:file_asset(rec,str(p),'20261009_校版预览/'+p.name,origin='local')
figure(rec,ASSETS/rec['folder']/'原始材料/20261009_校版预览/pathway_selected_dense_comparison.png','33 条入选通路并排比较。左侧 raw 色标采用实际观测范围，下界不为零；右侧仅为槽内排序。完整 329 通路矢量图在材料目录。')
folder_asset(rec,'paper/figures/v313_slotspe_20261007_v2/blca_a9kp_transport')

rec=add_record(13,'队列通路','BLCA 和 KIRC 队列通路权重差值','已有预测的队列分析','两队列矩阵与图已保存','按预测风险组观察通路最终聚合权重的描述性差异。','折内验证风险 midrank 百分位分 Q1–Q4；先对每位患者八槽等权平均，再按通路名对齐。图显示组均值减四组均值的等权平均。','BLCA 380 人、入选 29 条通路；KIRC 488 人、入选 33 条通路。绝对组间偏离较小，颜色不同不能替代数值大小。','两队列使用不同色标范围。通路由验证 attention 选取，不能把图解释为表达量、因果风险贡献或性能提升的机制证明。',['paper/figures/v313_manuscript_integrated/manifest.json','paper/figures/v313_slotspe_20261007_v2/blca_cohort/cohort_top_pathways.json','paper/figures/v313_slotspe_20261007_v2/kirc_cohort/cohort_top_pathways.json'],['scripts/plot_v313_slotspe_style.py','scripts/prepare_v313_manuscript_panels.py','survot_rank/evidence/slotspe_style.py'])
rows=[]
for c in ['blca','kirc']:
    q=meta['cohort'][c]
    rows.append([c.upper(),q['source_n'],' / '.join(map(str,q['group_counts'])),q['selected_pathways'],f"{q['absolute_delta_limit']:.3e}"])
    folder_asset(rec,f'paper/figures/v313_slotspe_20261007_v2/{c}_cohort')
    for stem in [f'fig7_attention_difference_{c}',f'figS3_raw_attention_{c}']:
        for ext in ['png','pdf','svg']:
            p=file_asset(rec,f'paper/figures/v313_manuscript_integrated/{stem}.{ext}')
            if ext=='png' and stem.startswith('fig7'):figure(rec,p,f'{c.upper()} 的绝对组均值差图。请读取色标数值，不按颜色深浅跨队列比较。')
table(rec,'队列分组与差值范围',['队列','患者数','Q1 至 Q4 人数','入选通路','最大绝对差'],rows)
file_asset(rec,'paper/figures/v313_manuscript_integrated/manifest.json')

rec=add_record(14,'两队列成本','Full 与 Direct 两队列实测成本','已训练模型的 profile','汇总记录已保存 计时待复核','比较 BLCA/KIRC Full 与 Direct 的性能及评估前向成本。','RTX 5090、float32、batch=1、2048 patches、warmup=5、repeats=30。每折测首位验证病例；横轴延迟是各折重复前向中位数的宏平均。','Full/Direct 的两队列平均 C-index 为 0.7731/0.7530，峰值 allocated 显存同为 181.39 MiB；平均延迟为 488.66/504.19 ms。','Full 仅 4/10 折延迟更低。BLCA fold0 的 Direct 为 650.62 ms、Full 为 482.76 ms，明显影响平均值；不得称稳定更快。原始逐次计时仍在服务器。BLCA 五折与两队列图复用同一批预测。',['paper/figures/v313_tradeoff_real_20261008/tradeoff_blca_kirc/efficiency_tradeoff.json'],['scripts/assemble_v313_tradeoff.py','scripts/plot_v313_reference_panels.py','scripts/run_v313_additional_evidence.py','survot_rank/evidence/additional.py','survot_rank/evidence/reference_panels.py'])
table(rec,'两队列十折宏平均',['方法','C-index','峰值 allocated 显存 MiB','前向延迟 ms'],[[v['id'].replace('dct_v313_','').title(),f"{v['cindex']:.4f}",f"{v['memory_mib']:.2f}",f"{v['latency_ms']:.2f}"] for v in profile['points']])
full,direct=profile['points'];rows=[];faster=0
for f in full['folds']:
    d=next(v for v in direct['folds'] if (v['cancer'],v['fold'])==(f['cancer'],f['fold']))
    faster+=f['latency_ms']<d['latency_ms']
    rows.append([f['cancer'],f['fold'],f"{f['latency_ms']:.2f}",f"{d['latency_ms']:.2f}",f"{d['latency_ms']-f['latency_ms']:+.2f}"])
assert faster==4
table(rec,'逐折前向延迟',['队列','Fold','Full ms','Direct ms','Direct 减 Full ms'],rows)
files=folder_asset(rec,'paper/figures/v313_tradeoff_real_20261008/tradeoff_blca_kirc')
figure(rec,files['efficiency_tradeoff.png'],'原始两均值成本图，长方法名和版式待修。结论以本条逐折表为准，平均延迟差不代表稳定加速。')
for ext in ['png','pdf','svg','json']:file_asset(rec,f'paper/figures/v313_tradeoff_real_20261008/efficiency_tradeoff.{ext}','BLCA_五折子集/efficiency_tradeoff.'+ext)

rec=add_record(15,'十队列成本','Full 十队列实测成本分布','完整模型的跨队列 profile','50 折汇总记录已保存','记录固定输入规模下 Full 在十队列的评估前向开销。','与 E014 同 GPU、精度和输入规模；每队列五折宏平均，仅 Full，不含 Direct 或外部方法。','十队列峰值 allocated 显存约 180.3–181.4 MiB，延迟约 479.5–499.5 ms，C-index 约 0.6300–0.8224。','这是十队列单方法成本分布，不证明优于外部方法。前向范围可能包含辅助重建，不等于完整病理预处理与部署流程；不同癌种的成绩差不能直接作难度因果解释。',['paper/figures/v313_tradeoff_real_20261008/tradeoff_ten_cohort_full_only/efficiency_per_cohort.json'],['scripts/plot_v313_tradeoff_per_cohort.py','scripts/run_v313_additional_evidence.py','survot_rank/evidence/additional.py'])
table(rec,'十队列 Full profile',['队列','C-index','峰值 allocated 显存 MiB','延迟 ms','参数量 M'],[[v['cancer'],f"{v['cindex']:.4f}",f"{v['memory_mib']:.2f}",f"{v['latency_ms']:.2f}",f"{v['parameters']/1e6:.3f}"] for v in profile10['points']])
files=folder_asset(rec,'paper/figures/v313_tradeoff_real_20261008/tradeoff_ten_cohort_full_only')
figure(rec,files['efficiency_per_cohort.png'],'同一协议下的十队列 Full 分布，每个点是一队列五折汇总。')

rec=add_record(16,'公开结果参照','外部方法公开报告值参照','文献结果整理','公开表已汇总 同条件复现未完成','保存现论文外部参照的来源、数值和比较边界。','21 个外部方法与 DCT；十队列公开均值按癌种名称对齐。不同论文的编码器、纳入患者、划分与 checkpoint 选择不同。','183 个可对齐的外部队列均值已汇总；公开参照可以帮助定位报告值，不作为当前同条件 SOTA 证据。','MCAT、MOTCat、CMTA、LD-CVAE 等外部方法同条件 checkpoint/profile 仍缺。ProtoPathway 的 OS 结果不混入 DSS 排名。',['paper/V313_OT_UNI2H_COMPARISON_INPUT.json'],['scripts/build_v313_comparison_table.py','scripts/build_v313_literature_comparison.py'])
for cancers in [['BLCA','BRCA','COADREAD','HNSC','KIRC'],['LUAD','LUSC','SKCM','STAD','UCEC']]:
    rows=[]
    for v in comparison['models']:
        means=v.get('reported_means',{})
        if v.get('id','').startswith('dct') or v.get('model','').startswith('DCT'):
            means={row[0]:float(row[-1].split('±')[0]) for row in main[1:]}
        rows.append([v.get('model',v.get('id',''))]+[f'{means[c]:.3f}' if c in means and means[c] is not None else 'NR' for c in cancers])
    table(rec,'公开参照均值 '+ ' '.join(cancers),['方法']+cancers,rows)
file_asset(rec,'paper/V313_OT_UNI2H_COMPARISON_INPUT.json')
folder_asset(rec,'paper/tables/v313_ot_uni2h_comparison')

print('Evidence records prepared:',len(records),flush=True)
# Freeze relevant source files and their resolvable local imports. This is an
# audit copy, not a new executable checkout or a relocated training project.
def git(*args):return subprocess.check_output(['git',*args],cwd=ROOT)
all_paths=set(git('ls-tree','-r','--name-only',COMMIT).decode('utf8').splitlines())
requested={MODEL,'survot_rank/research/methods/legacy/experimental/dct_v311_slot_interpretable/model.py','survot_rank/research/methods/catalog.py'}
for rec in records:requested.update(rec['scripts'])
requested.update(p for p in all_paths if p.startswith('configs/dct_v313_') and p.endswith('.yaml'))
queue=deque(sorted(requested));copied={};missing=[]
while queue:
    rel=queue.popleft()
    if rel in copied:continue
    if rel not in all_paths:
        missing.append(rel);continue
    data=git('show',f'{COMMIT}:{rel}')
    out=SNAP/rel;out.parent.mkdir(parents=True,exist_ok=True);out.write_bytes(data)
    copied[rel]=dict(path=rel,sha256=hashlib.sha256(data).hexdigest(),source_commit=COMMIT)
    if not rel.endswith('.py') or rel=='survot_rank/research/methods/catalog.py':continue
    try:tree=ast.parse(data.decode('utf-8-sig'))
    except (UnicodeDecodeError,SyntaxError):continue
    package=rel[:-3].split('/')[:-1]
    for node in ast.walk(tree):
        mods=[]
        if isinstance(node,ast.Import):mods=[v.name for v in node.names]
        elif isinstance(node,ast.ImportFrom):
            if node.level:base='.'.join(package[:len(package)-node.level+1]+([node.module] if node.module else []))
            else:base=node.module or ''
            mods=[base]+[base+'.'+v.name for v in node.names]
        for mod in mods:
            if not mod.startswith(('survot_rank','models','datasets','utils')):continue
            for candidate in [mod.replace('.','/')+'.py',mod.replace('.','/')+'/__init__.py']:
                if candidate in all_paths and candidate not in copied:queue.append(candidate)
code_manifest=dict(commit=COMMIT,files=list(copied.values()),unresolved_requested=sorted(set(missing)),purpose='Frozen source copy for provenance; datasets/checkpoints/environment remain in original project/server')
(SNAP/'代码来源清单.json').write_text(json.dumps(code_manifest,ensure_ascii=False,indent=2),encoding='utf8')
for rec in records:
    index=dict(experiment=rec['id'],title=rec['title'],source_commit=COMMIT,shared_model=MODEL,entries=[dict(path=p,present=p in copied,sha256=copied.get(p,{}).get('sha256')) for p in rec['scripts']],shared_snapshot='../共享代码快照_2b1c622',execution_note='仅作源代码溯源，整理阶段没有执行训练、推理或导出作业。以原工程环境与历史 overrides 复现。')
    (CODE/rec['folder']/'代码索引.json').write_text(json.dumps(index,ensure_ascii=False,indent=2),encoding='utf8')
    (CODE/rec['folder']/'README.md').write_text('# '+rec['id']+' '+rec['title']+'\n\n对应入口见代码索引.json。实际源文件位于 ../共享代码快照_2b1c622/，保留原工程相对路径。\n\n快照用于核对代码，不替代历史配置、checkpoint 或服务器运行环境。\n',encoding='utf8')
    (ASSETS/rec['folder']/'实验材料索引.json').write_text(json.dumps(rec,ensure_ascii=False,indent=2),encoding='utf8')
(MAINT/'实验目录_建档快照.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf8')
(MAINT/'材料来源清单.json').write_text(json.dumps(dict(source_commit=COMMIT,files=manifest),ensure_ascii=False,indent=2),encoding='utf8')
file_source=MAINT/'来源论文_2b1c622.md';shutil.copy2(SRC/'paper/DCT_v313_初稿.md',file_source)
print('Code snapshot files:',len(copied),'asset files:',len(manifest),flush=True)

# Editable Word main. Tables and figures are embedded; hyperlinks are relative
# to this document so the whole archive folder can be copied together.
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7)
sec.left_margin=sec.right_margin=Inches(.8)
sec.header_distance=sec.footer_distance=Inches(.3)
for style_name,size in [('Normal',11),('Title',28),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Caption',9.5)]:
    style=doc.styles[style_name];style.font.name='Calibri';style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if style_name in ['Normal','Caption'] else '微软雅黑')
    style.paragraph_format.space_after=Pt(5)
    style.paragraph_format.line_spacing=1.12
    if style_name in ['Heading 1','Heading 2']:style.paragraph_format.keep_with_next=True
normal=doc.styles['Normal'];normal.paragraph_format.widow_control=True
header=sec.header.paragraphs[0];header.text='DCT 实验事实主档';header.style=doc.styles['Caption']
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);footer.add_run(' 页')
settings=doc.settings.element;update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)

def p(text='',style=None):return doc.add_paragraph(text,style)
def heading(text,level=1,new=False):
    if new:doc.add_page_break()
    return doc.add_heading(text,level)
def link(par,text,target):
    h=OxmlElement('w:hyperlink');rid=par.part.relate_to(urllib.parse.quote(str(target).replace('\\','/'),safe='/:.'),RT.HYPERLINK,is_external=True);h.set(qn('r:id'),rid)
    run=OxmlElement('w:r');props=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'245778');props.append(color);under=OxmlElement('w:u');under.set(qn('w:val'),'single');props.append(under);run.append(props);t=OxmlElement('w:t');t.text=text;run.append(t);h.append(run);par._p.append(h)
def native_table(data):
    p(data['title'],'Caption');headers=data['headers'];rows=data['rows']
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if len(headers)==7:w=[.65]+[.85]*5+[2.0]
    elif len(headers)==6:w=[1.5]+[1.08]*5
    elif len(headers)==5:w=[1.0,.65,1.45,1.8,2.0]
    elif len(headers)==4:w=[1.2,1.6,2.2,1.9]
    else:w=[1.3,3.9,1.7] if len(headers)==3 else [6.9/len(headers)]*len(headers)
    w=np.array(w)*6.9/sum(w)
    for j,h in enumerate(headers):t.rows[0].cells[j].text=h
    for row in rows:
        for cell,value in zip(t.add_row().cells,row):cell.text=value
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(float(w[ci]));cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            pr.append(borders);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'264D66' if ri==0 else ('F1F5F8' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(0);par.paragraph_format.space_before=Pt(0);par.paragraph_format.line_spacing=1.03
                par.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs:
                    run.font.size=Pt(9.5 if len(headers)>=6 else 10);run.font.bold=ri==0
                    run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    p()
    return t

def picture(path,caption,max_h=6.6):
    with Image.open(path) as im:width,height=im.size
    ratio=height/width;w=min(6.9,max_h/ratio)
    par=p();par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pic=par.add_run().add_picture(path,width=Inches(w));pic._inline.docPr.set('descr',caption)
    p(caption,'Caption')

p('DCT 实验总档案','Title');p('实验名称 结果成果与材料索引','Subtitle')
p('本文件是 DCT 实验事实的长期主档。当前记录覆盖 v3.13 论文使用的训练目标、机制对照、预测分析、解释图和实测成本。以后新增实验在本文件续写，保留编号、设置、结果与来源。')
p('本次建档共 16 个实验与分析条目。训练结果、固定模型分析和外部论文参照分别标明，复用 Full 模型的材料不累计为独立训练次数。')
native_table(dict(title='档案信息',headers=['项目','内容'],rows=[['主档级别','实验事实主档 供论文与汇报引用核对'],['更新时间',DATE],['冻结来源','Git 2b1c622'],['主要协议','UNI2-h DSS 五折 seed 3 legacy_val 最佳验证 checkpoint'],['主要成果','七组目标消融 两组机制控制 十队列 Full 与 KM 通路解释 RTX 5090 成本'],['续写方式','复制末尾模板 新实验另起页 目录可更新'],['代码位置','实验代码 按编号索引 共享冻结源代码'],['结果位置','表格图片 每项实验独立目录']]))
p('主档的解释必须服从对应原始产物与实验协议。新汇报文本不自动覆盖旧实验事实；若发现数值或结论有误，应在原条目中登记修订并保留旧版本。')
p('服务器原始训练结果根目录为 /data1/DCT-Reg/results。当前材料目录保存已提交的表格、图和结果 JSON，未把模型 checkpoint 与整套数据迁移进档案。')
heading('目录',new=True)
par=p();f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');par._p.append(f)
p('目录与导航窗格使用各实验的一级标题。Word 中可右键目录选择更新整个目录。','Caption')
heading('档案使用与共同实验协议',new=True)
p('Word 是持续续写的主档。维护资料中的 JSON 和生成程序是本次建档快照，不能在手工续写后未经核对覆盖本文件。')
native_table(dict(title='共同开发协议',headers=['设置','当前记录'],rows=[['病理特征','UNI2-h 1536 维 每例最多 2048 patches'],['模态与目标','病理与通路组学 疾病特异性生存 DSS'],['划分与选择','患者级五折 train/validation 验证集同时用于 epoch 选择与评分'],['训练','seed 3 最多 30 epochs batch 8 AdamW 学习率与衰减 0.0005'],['模型','隐维 256 两模态各 8 槽 3 次迭代 4 阶段 3 几何'],['汇总','每队列五折均值与样本标准差 十队列按癌种等权'],['测量','RTX 5090 float32 batch 1 2048 patches warmup 5 repeats 30'],['证据边界','最佳验证结果不是独立外层测试 折间标准差不是多 seed 方差']]))
p('材料按 E001 至 E016 的固定编号对应。每次新增需记录名称、目的、版本、数据和终点、划分、seed、配置与 overrides、checkpoint 选择、指标、代码、产物、结论和待核对项。失败或未完成的实验也保留状态，不补填猜测数值。')
p('引用顺序为原始产物与身份/配置核对、主档实验条目、论文与汇报摘要。发现不一致先核对来源，记录修订后再用于论文。公开值仅作参照，不冒充本人同条件复现。')
p('代码文件保留在原工程。实验代码目录提供按冻结提交提取的源文件及索引，便于追溯；数据集、环境和历史 overrides 仍需按原运行位置核对。')

for rec in records:
    heading(rec['id']+' '+rec['title'],new=True)
    p('类型：'+rec['kind']+'    状态：'+rec['status'])
    p('目的：'+rec['purpose']);p('设置：'+rec['setting'])
    # Dedicated pages for large secondary tables keep each experiment legible.
    if rec['id']=='E007':
        for t in rec['tables'][:2]:native_table(t)
        p('主要成果：'+rec['conclusion']);p('待核对：'+rec['pending'])
        heading('E007 十队列结果',level=2,new=True)
        for t in rec['tables'][2:]:native_table(t)
    elif rec['id']=='E016':
        p('主要成果：'+rec['conclusion']);p('待核对：'+rec['pending'])
        for t in rec['tables']:
            heading(t['title'],level=2,new=True);native_table(t)
    else:
        for t in rec['tables']:native_table(t)
        if rec.get('note'):p(rec['note'],'Caption')
        p('主要成果：'+rec['conclusion']);p('待核对：'+rec['pending'])
    if rec['id']=='E007':p('待核对：'+rec['pending'])
    par=p('相关材料：');link(par,'结果表格与图片目录',str((ASSETS/rec['folder']).relative_to(DEST)))
    par.add_run('    ');link(par,'实验代码索引',str((CODE/rec['folder']/'代码索引.json').relative_to(DEST)))
    p('来源：'+ '；'.join(rec['sources'])+'。冻结提交 2b1c622。','Caption')
    # Pictures are readable exhibits rather than every archived duplicate.
    if rec['figures']:
        heading(rec['id']+' 图像成果',level=2,new=True)
        max_h=3.2 if rec['id']=='E013' else 6.4
        for item in rec['figures']:picture(item['path'],item['caption'],max_h=max_h)

heading('待补项与修订记录',new=True)
native_table(dict(title='当前待补清单',headers=['条目','待补内容','当前状态'],rows=[['E005 E006','与 Full 单支系数匹配的 self-only/cross-only 重训','未作为完成证据'],['E008 E009','Full 与控制臂患者 结局 配置的最终配对核对','审计控制臂已保存'],['E010','患者级最佳预测对齐 风险差分 重建误差 边缘残差','十折汇总存在'],['E011','总览脚注 UCEC 图例 尾部风险人数检查','KM 数据齐全 图待校版'],['E012','真实组织图 Top-5 组织块 KIRC 病例 跨病例槽诊断','非组织病例分析已有'],['E014','BLCA fold0 计时复核 逐次测量波动 长标签排版','不能称稳定加速'],['E016','外部方法同条件 checkpoint 与测量','公开表不能替代复现'],['历史补录','v3.10 v3.11 v3.15 v3.16 分版本登记','不并入本版 v3.13 数值']]))
native_table(dict(title='修订日志',headers=['日期','修订内容','依据'],rows=[[DATE,'建立主档与编号目录 归档代码和表格图片','冻结提交 2b1c622'],[DATE,'记录成本均值受单折影响 通路 raw 与 percentile 的显示边界','逐折 profile JSON 与原始 329×8 权重']]))
heading('新实验续写模板',new=True)
p('新增实验时复制本页，放在待补项之前，改成下一个未使用编号。标题使用一级标题并保留另起页；完成后更新目录。首次新增可从 E017 开始。')
p('E017 实验名称','Heading 2')
p('类型：训练实验／固定模型分析／数据分析／文献参照    状态：计划中／进行中／失败／完成')
p('实验目的：')
p('版本与提交：    运行日期：    负责记录人：')
p('数据与终点：    特征编码器：    纳入样本与患者隔离：')
p('划分与 seed：    epoch 与 checkpoint 选择：')
p('配置文件与实际 overrides：')
p('有效目标系数与对照差异：')
native_table(dict(title='结果表',headers=['队列或条件','指标与重复单位','结果与离散度'],rows=[['待填写','待填写','未测量'],['待填写','待填写','未测量']]))
p('图片与图注：有图时插入，并说明颜色、单位、样本和选择规则。')
p('相关代码：实验代码／E017_实验名称／代码索引.json')
p('表格图片：表格图片／E017_实验名称／')
p('成果与可支持结论：')
p('限制与不能支持的结论：')
p('原始产物 来源路径 SHA256 与配对核对状态：')
p('与既有实验的复用关系：')
p('后续动作与修订记录：')
DOCX.parent.mkdir(parents=True,exist_ok=True);doc.save(DOCX)
readme='''# DCT 实验档案\n\n主档为 DCT实验总档案.docx。以后新增实验直接在 Word 主档续写，复制末尾模板并更新目录。\n\n- 实验代码：每项实验一个编号目录，代码索引指向共享代码快照_2b1c622。原工程文件没有移动。\n- 表格图片：每项实验一个编号目录，含可编辑 CSV 结果表、原始图/JSON 和材料索引。\n- 维护资料：本次建档的来源清单和生成快照。手工续写 Word 后不得直接用旧生成程序覆盖主档。\n\n新增步骤：分配下一个编号；在 Word 复制模板另起页；新增对应的代码与表格图片目录；填写实际结果与原始来源；更新目录和修订日志。已有编号不复用，原结果被推翻时保留修订原因。\n\n主档用于实验事实优先查阅；若与原始产物不一致，以可核验产物和协议纠正主档，不以新的汇报口号覆盖旧结果。\n'''
(DEST/'README.md').write_text(readme,encoding='utf8')
entry=ROOT/'EXPERIMENTS.md'
if not entry.exists():entry.write_text('# DCT 实验事实主档入口\n\n长期续写的主档为 [DCT实验总档案.docx](实验档案/DCT实验总档案.docx)。\n\n论文与汇报中的实验事实应先核对主档及对应原始产物。代码索引在 实验档案/实验代码/，表格图片在 实验档案/表格图片/。旧总结不自动覆盖主档，数据冲突以可核验原始产物和真实协议为准。\n',encoding='utf8')
shutil.copy2(__file__,MAINT/'build_archive_初次建档.py')
check=dict(document=str(DOCX),experiments=len(records),tables=len(doc.tables),images=len(doc.inline_shapes),source_commit=COMMIT,source_files=len(manifest),code_files=len(copied),unresolved_code_requests=sorted(set(missing)))
(MAINT/'建档核验.json').write_text(json.dumps(check,indent=2,ensure_ascii=False),encoding='utf8')
print(json.dumps(check,indent=2,ensure_ascii=False),flush=True)
