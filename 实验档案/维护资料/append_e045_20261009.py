from pathlib import Path
import copy,hashlib,json,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
MASTER=ROOT/'实验档案/DCT实验总档案.docx'
BACKUP=ROOT/'历史备份/实验档案/20261009_E045_续写前/DCT实验总档案.docx'
FIG=ROOT/'实验档案/表格图片/E045_模块作用核验/现有记录_20261009'

def checksum(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def paragraph(d,text,style=None):
    p=d.add_paragraph(text,style=style)
    for run in p.runs:
        run.font.name='Calibri';run.font.size=Pt(10.5 if style is None else 9.5)
        run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    return p

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers));t.style='Table Grid';t.autofit=False
    t._tbl.remove(t._tbl.tblPr);t._tbl.insert(0,copy.deepcopy(d.tables[0]._tbl.tblPr))
    for i,(name,width) in enumerate(zip(headers,widths)):t.columns[i].width=Inches(width);t.rows[0].cells[i].text=name
    for values in rows:
        cells=t.add_row().cells
        for c,text in zip(cells,values):c.text=str(text)
    for row_index,row in enumerate(t.rows):
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:
                    run.font.size=Pt(9);run.font.name='Calibri';run.bold=row_index==0
                    run.font.color.rgb=RGBColor.from_string('FFFFFF' if row_index==0 else '111111')
                    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'254B61' if row_index==0 else ('F1F5F8' if row_index%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
    return t

if BACKUP.exists():raise RuntimeError('Backup already exists; do not overwrite or repeat authoring')
original=Document(MASTER)
if any(p.style.name=='Heading 1' and p.text.startswith('E045') for p in original.paragraphs):raise RuntimeError('E045 already registered')
BACKUP.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(MASTER,BACKUP)
old_hash=checksum(BACKUP)
original_paragraph_count=len(original.paragraphs)
d=original
for p in d.paragraphs:
    if p.text=='当前保留 19 个 v3.13 相关实验与分析条目。复用 Full 模型的图表不累计为独立训练次数；公开文献参照不计作本人训练实验。':
        p.text='当前保留 20 个 v3.13 相关实验与分析条目，新增 E045 为既有模块对照再分析及待运行可视化入口。复用 Full 模型的图表不累计为独立训练次数；公开文献参照不计作本人训练实验。'
    if p.text.startswith('本轮按用户明确要求，将主档范围收窄到 v3.13：'):
        p.add_run(' 2026-10-09 E045 已续写，当前共 20 条，下一编号 E046。')
    if '首次新增可从 E045 开始' in p.text:
        for r in p.runs:r.text=r.text.replace('首次新增可从','下一项新增从').replace('E045','E046')
    if 'E045 实验名称'==p.text:p.text='E046 实验名称'
    if 'E045_实验名称' in p.text:p.text=p.text.replace('E045_实验名称','E046_实验名称')

d.add_page_break();d.add_heading('E045 v313 模块作用核验与可视化',level=1)
paragraph(d,'状态：两张既有记录图已生成；真实患者的运输重放三联图、槽诊断图与修正版配对检索尚待服务器执行。本条是再分析与代码登记，不是新增训练结果。')
d.add_heading('核心 idea 与模块归属',level=2)
paragraph(d,'v3.13 让 cross 重建复用预测使用的事实 OT 对应：先把病理槽映射到组学槽坐标，再用通路查询解码器恢复通路 token；self 重建约束组学槽保留信息。目标是通过重建监督改善跨模态对应和表征，而非只靠多一个解码器或高质量特征得分。')
table(d,['组成','实际作用','归属'],[
['语义槽与共享坐标','组织两种模态的槽表征','继承的框架'],
['分阶段多几何 OT','计算槽之间的跨模态对应并供预测使用','继承的框架'],
['运输感知 self/cross 重建','cross 复用事实计划，解码通路 token；self 保留组学信息','v3.13 主要新增']],[1.6,3.3,1.8])
paragraph(d,'重建目标是已观测组学编码得到的潜在 token，目标 detach；阶段融合 gate 也 detach。实际组学参与计划计算，不应把这一任务称为仅病理预测未观测基因。梯度可达证明实现接通，真实模块收益仍需控制实验。')
d.add_heading('已有支持和未完成的证明',level=2)
table(d,['对照','BLCA Full 提升','KIRC Full 提升'],[
['Direct 仅改 cross 输入','+0.0209，3/5 折','+0.0193，4/5 折'],
['Independent 改预测与重建计划','+0.0146，3/5 折','+0.0170，5/5 折']],[2.7,2,2])
paragraph(d,'这些 seed3、legacy_val 记录支持运输重建配方带来的训练收益，比 Full 单独分数更接近模块验证；但完整患者与配置配对、多 seed、独立测试仍未齐备。Independent 同时改两个位置，不能据此单独归因 cross。Exp4/5 历史权重 0.025 与 Full 各 0.05 不同，严格删除实验应另立。')
paragraph(d,'固定 Full 替换同边际计划后 BLCA 排序不变，KIRC 仅一折略变。这不支持排序强依赖精细对应；也不能由排序不变推断全部风险与重建不变。')

d.add_page_break();d.add_heading('E045 现有真实记录图',level=2)
d.add_picture(str(FIG/'training_module_gains.png'),width=Inches(6.7))
paragraph(d,'图 E045a Full−训练控制的逐折 C-index 差值。正值支持 Full，负值明确保留。Full 四位小数历史折值和控制原值均保留；未添加显著性标记。',style='Caption')
d.add_picture(str(FIG/'frozen_plan_ranking_response.png'),width=Inches(6.7))
paragraph(d,'图 E045b 固定 Full 的计划替换响应。0 为事实计划，1 为同事实边际乘积；KIRC fold2 变化约 +0.00068，其余曲线为零。历史图缺患者级重放核对，不能补成强机制证明。',style='Caption')
paragraph(d,'代码：实验代码/E045_模块作用核验/visualize_v313_idea.py。图表：表格图片/E045_模块作用核验/现有记录_20261009。13 个来源及产物哈希见 evidence.json；表格见 paired_fold_gains.csv。')

d.add_page_break();d.add_heading('E045 新可视化实验判据与执行入口',level=2)
paragraph(d,'优先复用现有 checkpoint 和经过严格校验的 export，不需要新增训练。若旧 export 只有 alpha=0 或只有病例图，需同 checkpoint 重放后再作图；由用户在服务器执行。')
table(d,['诊断','图中要看的真实变化','何时不支持假说'],[
['运输干预三联图','同患者逐折的排序、风险/IQR、cross token 重建误差','计划被破坏但重建误差不升；或患者/边际不匹配'],
['错配与去中心重建','native 对错配 WSI、训练均值；使用检索 v2','native 未优于控制；旧 v1 结果不得替代'],
['槽功能三联图','原始表征余弦、跨槽 hazard 方差、通路 attention 熵','不能凭漂亮 percentile 或单一统计认定生物学分工']],[1.4,2.8,2.5])
paragraph(d,'replay 绘图器检查五折完整性、患者结局、最佳预测重放、原始来源哈希、通路顺序、alpha=0 与运输边际残差；不满足就停止。跨 checkpoint 的原始风险不混排名，风险变化按各折事实风险 IQR 归一化。')
paragraph(d,'服务器步骤：实验代码/E045_模块作用核验/SERVER_PROMPT.txt；真实导出清单结构：replay_input.example.json。新的原始 NPZ 保存在服务器 results；轻量图与汇总另放本条图表目录的新子目录，不覆盖旧材料。')
paragraph(d,'7 项代码校验通过，包括排序不变但风险变化的合成反例、患者与源指标不匹配、运输边际变化、非归一化 attention、零风险 IQR 的拒绝。这是代码校验，不是模块效果或真实患者诊断结果。')
paragraph(d,'九行十队列五折表共 450 条，现有 130 条，尚未覆盖 320 条；先查服务器已有材料，不重复训练 Full。验证 diversity 的单独作用需 Exp2/Exp3 同条件对照，Full 的槽诊断只能描述当前状态。')
paragraph(d,'结论：模块收益已有初步支持，精细跨模态对应的机制尚未闭环。重建曲线、错配响应与槽诊断都需真实结果才能升级结论；允许结果不支持原假说，不预设成功。')
assert checksum(BACKUP)==old_hash
d.save(MASTER)
receipt={'date':'2026-10-09','experiment':'E045','backup':str(BACKUP.relative_to(ROOT)).replace('\\','/'),'backup_sha256':old_hash,'master_sha256_before':old_hash,'master_sha256_after_authoring':checksum(MASTER),'records':20,'new_training_runs':0,'pending_real_replay':True,'visual_qa':'pending render','old_paragraphs':original_paragraph_count}
(ROOT/'实验档案/维护资料/E045_续写核验.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8',newline='\n')
print('Master backed up and E045 appended; original experiment material untouched.')
