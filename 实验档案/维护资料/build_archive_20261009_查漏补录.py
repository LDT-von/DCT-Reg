from pathlib import Path
import ast,csv,hashlib,json,re,shutil,subprocess,urllib.parse,sys
import numpy as np
from copy import deepcopy
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image
sys.stdout.reconfigure(encoding='utf8')
ROOT=Path(r'E:\DCT-Reg');BASE=ROOT/'实验档案';TASK=ROOT/'.review/experiment-ledger-audit-20261009';SRC=TASK/'source_snapshot';FIGSRC=ROOT/'.review/experiment-ledger-20261009/source_snapshot';DOC=BASE/'DCT实验总档案.docx';MAINT=BASE/'维护资料';CODE=BASE/'实验代码';ASSETS=BASE/'表格图片';COMMIT='2b1c622f418c619dc20f973ddeb839ff12d30d88'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
backup=MAINT/'历史版本/20261009_补录前'
backup.mkdir(parents=True,exist_ok=True)
if (backup/DOC.name).exists():assert sha(DOC)==sha(backup/DOC.name),'Main document has changed since backup; preserve it.'
for p in [DOC,MAINT/'实验目录_建档快照.json',MAINT/'材料来源清单.json',MAINT/'建档核验.json']:
 if not (backup/p.name).exists():shutil.copy2(p,backup/p.name)
# Capture the existing source state before this revision. No source relocation.
protected={}
for folder in ['results','results_fixed_anchors','audit_results','experiments','paper_outputs','paper','docs','scripts','configs','survot_rank','standalone']:
 for p in (ROOT/folder).rglob('*'):
  if p.is_file() and '__pycache__' not in p.parts:protected[str(p.relative_to(ROOT))]=sha(p)
for p in ROOT.glob('*.md'):protected[str(p.relative_to(ROOT))]=sha(p)
(TASK/'source_before.json').write_text(json.dumps(protected,ensure_ascii=False,indent=2),encoding='utf8')
records=json.loads((MAINT/'实验目录_建档快照.json').read_text(encoding='utf8'));new=[];copied=[]
def locate(rel):
 for root in [SRC,FIGSRC,ROOT]:
  if (root/rel).is_file():return root/rel
 raise FileNotFoundError(rel)
def j(rel):return json.loads(locate(rel).read_text(encoding='utf8'))
def mt(rel):
 tables=[];block=[]
 for l in locate(rel).read_text(encoding='utf8').splitlines()+['']:
  if l.startswith('|'):block.append([x.replace('**','').strip() for x in l.strip().strip('|').split('|')])
  elif block:
   if len(block)>2:tables.append(dict(title='原报告数值',headers=block[0],rows=block[2:]))
   block=[]
 return tables

def record(num,slug,title,status,setting,conclusion,sources,scripts=[],pending='历史运行配置、划分、checkpoint 和患者身份尚需核对；数值不合并进 v3.13 主表。'):
 r=dict(id=f'E{num:03d}',folder=f'E{num:03d}_{slug}',title=title,kind='补录实验或审计',status=status,purpose='保留该项实验的名称、结果与证据状态。',setting=setting,conclusion=conclusion,pending=pending,sources=sources,scripts=scripts,tables=[],figures=[],assets=[])
 (ASSETS/r['folder']).mkdir(parents=True,exist_ok=True);(CODE/r['folder']).mkdir(parents=True,exist_ok=True)
 for rel in sources:
  try:p=locate(rel)
  except FileNotFoundError:continue
  out=ASSETS/r['folder']/'原始材料'/rel;out.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,out)
  entry=dict(experiment=r['id'],source=rel,source_kind='git_snapshot' if p.is_relative_to(SRC) or p.is_relative_to(FIGSRC) else 'local_worktree_snapshot',source_commit=COMMIT if p.is_relative_to(SRC) or p.is_relative_to(FIGSRC) else None,target=str(out.relative_to(BASE)),sha256=sha(out))
  copied.append(entry);r['assets'].append(entry)
 new.append(r);return r

def table(r,title,headers,rows):
 r['tables'].append(dict(title=title,headers=headers,rows=[[str(x) for x in row] for row in rows]))
def parsed(r,rel,index=0,limit=None):
 t=mt(rel)[index];table(r,t['title'],t['headers'],t['rows'][:limit] if limit else t['rows'])
def folder_sources(folder):
 return [str(p.relative_to(SRC)).replace('\\','/') for p in (SRC/folder).rglob('*') if p.is_file()]
def fig_sources(folder):
 return [str(p.relative_to(FIGSRC)).replace('\\','/') for p in (FIGSRC/folder).rglob('*') if p.is_file()]
ARCH='scripts/_archive_v310_v311/'
r=record(17,'v310目标消融','v3.10 BLCA 四臂目标消融','历史报告有五折表 原始训练日志未同步','Full、Direction-only、IPCW-only、NLL-only；旧报告最佳验证口径。','报告 Full 为 0.7175，Direction 为 0.7087。仅凭均值不能证明方向正确或目标协同；不能替代 E001 至 E007。',['EXPERIMENT_DATA_SUMMARY.md','DCT_v310_Ablation_Results_COMPLETE.md','VERIFIED_RESULTS_2026_09_05.md'],[ARCH+'run_dct_v310_experiments.py'])
t=mt('EXPERIMENT_DATA_SUMMARY.md')[0];table(r,'历史五折目标消融',['变体','均值','标准差'],[[x[0],x[6],x[7]] for x in t['rows']])
r=record(18,'v310机制对照','v3.10 锚点和耦合机制对照','历史报告部分折完成','BLCA、LUSC、UCEC；cross-fold frozen、permuted、fixed coupling、noisy anchors。','各控制臂折数与 Full 对照范围不同，历史汇总不能直接证明运输机制必要性。Stage jitter 未找到分数。',['EXPERIMENT_DATA_SUMMARY.md','DCT_v310_Experiments_Report.md'],[ARCH+'run_dct_v310_experiments.py'])
t=mt('EXPERIMENT_DATA_SUMMARY.md')[1];table(r,'历史机制控制覆盖',t['headers'][:5],[x[:5] for x in t['rows']])
r=record(19,'v310跨癌','v3.10 五癌种历史主结果与 KIRC 补折','历史报告记录 25 折 外层证据未完成','旧 UNI 及 50 epoch 配方；不是 v3.13 UNI2-h 十队列。','KIRC 五折为 0.8270，旧两折 0.8579 已作废。五癌种均值约 0.6964；原登记外层测试仍 pending。',['experiments/RESULTS.md','CRITICAL_EXPERIMENTS_MISMATCH.md'],[ARCH+'run_dct_v310_final_cross_cancer.py',ARCH+'kirc_refill_5fold.py'])
t=next(t for t in mt('experiments/RESULTS.md') if t['headers'][0]=='癌种' and '折数' in t['headers']);table(r,'五癌种历史报告',t['headers'],t['rows'])
r=record(20,'TGSR反馈','v3.2 TGSR 四臂反馈消融','历史报告记录 20 个训练任务','BLCA 五折；无反馈、self update、attention feedback、OT feedback，报告 NLL-only。','OT feedback 报告均值 0.7057，baseline 为 0.6828；不能由此证明锚点方向，也不能与不同配方 v3.10 直接归因比较。',['V32_TGSR_EXPERIMENT_RESULTS.md'],[ARCH+'run_dct_v32_experiments.py'])
t=mt('V32_TGSR_EXPERIMENT_RESULTS.md')[0];table(r,'历史反馈消融均值',['变体','报告均值','报告标准差'],[[x[1],x[2],x[3]] for x in t['rows']])
r=record(21,'TGSR目标','TGSR 目标函数与可学习反馈强度','历史报告五组五折 参考臂仅一折','NLL、IPCW、direction、full、full learned；tgsr_nll 的 fold0 预算与其他折不同。','Full 和 learned 报告均值为 0.7087/0.7084。完整参考臂未找到，不称为公平超越；保留协议差异。',['TGSR_BEST_EPOCH_RESULTS.md','TGSR_FINAL_VERDICT.md'],[ARCH+'run_dct_v32_objective_experiments.py'])
t=mt('TGSR_BEST_EPOCH_RESULTS.md')[0];table(r,'目标函数历史汇总',['变体','五折报告均值与标准差','覆盖状态'],[[x[1],x[-2],x[-1]] for x in t['rows']])
r=record(22,'锚点诊断','锚点跨折一致性和风险分离诊断','有历史报告 原始诊断 JSON 在服务器未同步','报告三项：跨折一致性、模态内距离、高低风险原型分离。','报告 WSI/omic 跨折相似度 0.694/0.798、归一化风险分离 0.0174。阈值与因果解释未独立复核，仅保留诊断结果。',['ANCHOR_EXPERIMENTS_REPORT_20260907.md','ANCHOR_DIAGNOSIS_SUMMARY_CN.md'],[ARCH+'anchor_consistency_analysis.py',ARCH+'anchor_separation_analysis.py',ARCH+'run_anchor_experiments.py'])
parsed(r,'ANCHOR_EXPERIMENTS_REPORT_20260907.md',0)
r=record(23,'固定锚点训练','固定锚点 BLCA 五折训练','历史报告有训练结果','报告使用预提取固定锚点，30 epochs。历史模型描述与实际前向是否一致需核对。','五折分数 0.7100、0.6684、0.6773、0.6842、0.7304，均值约 0.6941；不由此宣称方向机制通过。',['COMPLETE_STATUS_SUMMARY.md','FIXED_ANCHORS_EXECUTION_SUMMARY.md','README_FIXED_ANCHORS_EXPERIMENT.md'],[ARCH+'run_fixed_anchor_experiments.py',ARCH+'train_fixed_5fold.sh'])
parsed(r,'COMPLETE_STATUS_SUMMARY.md',0)
r=record(24,'固定锚点审计','固定锚点五折风险响应审计','五折 JSON 与四折 CSV 已保存','alpha 扫描；原始 fold0 仅摘要，fold1至4 CSV。','更新后的摘要按响应容差 1e-8 记录所有固定锚点折风险变化为零、响应率为零。不能沿用旧报告双方向 100% 的非严格单调判据。',folder_sources('results_fixed_anchors'),[ARCH+'audit_fixed_anchors_batch.py',ARCH+'e4_audit_adapted.py'])
table(r,'固定锚点响应摘要',['Fold','患者数','Low 响应率','High 响应率','Low 平均风险差'],[[f,j(f'results_fixed_anchors/audit_fixed_fold{f}_summary.json')['metrics']['n_patients'],0,0,0] for f in range(5)])
r['pending']='fold0 原始 CSV 未同步；保存的旧 CSV 中 true_event 出现非 0/1 数值，字段对应关系需复核，不能直接用于生存分析。'+r['pending']
r=record(25,'v311固定E4','v3.11 Fixed 五折嵌入干预审计','逐患者 CSV 和折级 JSON 已保存','使用 stage embedding 构建干预；模型内部嵌入操作。','五折 mr_mean 均值约 0.6129。记录该实现返回的指标，不把 0.5 自动解释为随机基线，也不作临床因果或显著性宣称。',folder_sources('results/e4_v311_fixed_audit')+['results/audit_comparison/v311_fixed_vs_v310_intervention_report.md'],[ARCH+'e4_v311_fixed_audit.py',ARCH+'make_e4_audit_report.py'])
vals=[j(f'results/e4_v311_fixed_audit/e4_v311_fixed_fold{f}.json') for f in range(5)];table(r,'逐折嵌入干预摘要',['Fold','n','mr_low','mr_high','mr_mean'],[[v['fold'],v['n_samples'],f"{v['mr_low']:.4f}",f"{v['mr_high']:.4f}",f"{v['mr_mean']:.4f}"] for v in vals])
r=record(26,'ProofA','Proof A 版本和特征配方训练记录','结构化 JSON 有五折与短程记录','v3.11 UNI 完整五折、UNI2-h 一轮 smoke，以及 v3.10 UNI2-h 不完整运行。','v3.11 UNI 五折最佳均值 0.7174；其余仅 smoke 或短程，不能作统一预算的版本优劣判断。',['results/proof_experiment_A.json'],[ARCH+'proof_experiments/proof_A_recipe_compare.py'])
a=j('results/proof_experiment_A.json')['by_recipe'];table(r,'配方覆盖和历史最佳值',['配方','折数','训练轮数范围','报告最佳均值'],[[k,len(v['folds']),f"{min(x['n_epochs'] for x in v['folds'].values())}–{max(x['n_epochs'] for x in v['folds'].values())}" if isinstance(v['folds'],dict) else '见 JSON',v['mean_best']] for k,v in a.items()])
r=record(27,'ProofB','Proof B 分模态槽方差诊断','五折 JSON 已保存','按 WSI 和 omic 分别计算槽方差；历史阈值区间 0.005 至 0.05。','WSI 五折均未达到报告定义的区间，omic 两折达到；不把最终通路权重熵当作相同槽方差指标。',['results/proof_experiment_B.json'],[ARCH+'proof_experiments/proof_B_variance_constraint_FIXED.py'])
b=j('results/proof_experiment_B.json');table(r,'逐折槽方差',['Fold','WSI 平均方差','Omic 平均方差','两模态通过'],[[v['fold'],v['wsi']['mean_var'],v['omic']['mean_var'],v['both_pass']] for v in b['folds']])
r=record(28,'ProofC','Proof C IPCW 训练轨迹','五折及短程比较 JSON 已保存','比较前五轮和后五轮 IPCW 目标；同时记录最佳验证 C-index。','v3.11 目标下降报告均值 0.2761，最佳验证均值 0.7174。目标下降不等于相对对照的性能提升。',['results/proof_experiment_C.json'],[ARCH+'proof_experiments/proof_C_ipcw_rank.py'])
c=j('results/proof_experiment_C.json');table(r,'v3.11 IPCW 目标轨迹',['Fold','前五轮均值','后五轮均值','下降','最佳 val C-index'],[[v['fold'],v['ipcw_first5'],v['ipcw_last5'],v['ipcw_drop'],v['val_c_best']] for v in c['v311']])
r=record(29,'ProofD','Proof D 槽风险排序与 KM 诊断','JSON 与报告已保存 部分原图本地缺失','五折病例槽风险排序统计和风险分层；旧产物病例数 380。','历史汇总 WSI/omic 相关约 -0.67/0.345，五折仅一折 KM p 小于 0.05；不混入新版十队列 KM。',folder_sources('results/proof_D_visualization'),[ARCH+'proof_D_visualization.py'])
d=j('results/proof_D_visualization/proof_D_metrics.json');table(r,'历史折级诊断',['Fold','n','WSI rho 均值','Omic rho 均值','KM p'],[[v['fold'],v['n'],v['mean_rho_wsi'],v['mean_rho_omic'],v['km_pval']] for v in d['summary']['combined_km']])
r=record(30,'v311槽比较','v3.11 Fixed 和 Unfixed 槽输出比较','历史报告 原始 hazard pkl 未同步','比较单槽输出、方差及模态级分数。','历史报告记录 Fixed/Unfixed 的 WSI 分数 0.5746/0.5548，omic 0.6944/0.6546。报告用曲线形状推断输出含义的解释未核实，不作为当前模型定义。',folder_sources('results/v311_fixed_proof'),[ARCH+'compare_v311_fixed_vs_unfixed.py',ARCH+'analyze_v311_fixed_proof.py'])
parsed(r,'results/v311_fixed_proof/v311_fixed_vs_unfixed_comparison.md',3)
r=record(31,'旧方向与零假设','旧版跨癌方向审计和运输零假设','JSON 已保存 审计判据需复核','跨癌 DCR、TV 和 BLCA factual、uniform、shuffled、anchor-swap；部分队列仅两折。','旧 uniform/shuffled 与 factual 的 DCR 完全相同；另有报告指出读取 factual 字段的问题。保留数值，但不能据此确认运输无效或已有效。',['paper_outputs/_archive_v310_v311/all_cancer_audits.json','paper_outputs/_archive_v310_v311/null_hypothesis_summary.json','paper_outputs/_archive_v310_v311/figure3_dose_response_partial.json','paper_outputs/_archive_v310_v311/DCR_AUDIT_REVISED.md','proof_experiments_validation_report.md','AUDIT_BEST_EPOCH_RESULTS.md'],[ARCH+'audit_dct_reg.py',ARCH+'generate_figure3_all_cancers.py'])
a=j(r['sources'][0]);table(r,'旧跨癌 DCR 汇总',['队列','折数','平均 DCR','平均 TV'],[[k.upper(),v['n_folds'],f"{v['mean_dcr']:.4f}",str(round(v.get('mean_tv',0),4)) if 'mean_tv' in v else '未记录'] for k,v in a.items()]);a=j(r['sources'][1])['summary_5fold'];table(r,'旧零假设摘要',['控制','平均 DCR','标准差'],[[k,f"{v['mean']:.4f}",f"{v['std']:.4f}"] for k,v in a.items()])
r=record(32,'旧E4分布','Direction 和 IPCW 旧 E4 风险分布','仅历史报告 原始 CSV 未找到','两种配置各五折，报告风险标准差和锚点距离。','风险 std 为 0.4526/0.7624；风险分布更窄不等于方向更正确，负风险分数本身也不证明预后改善。',['E4_AUDIT_FINAL_REPORT.md','ANALYSIS_EXECUTION_SUMMARY.md'],[ARCH+'run_all_e4_experiments.py',ARCH+'analyze_e4_results.py'])
table(r,'历史报告分布摘要',['指标','Direction','IPCW'],[['平均锚点距离','2.67 ± 1.52','2.48 ± 1.17'],['平均风险','-2.97 ± 0.55','-2.85 ± 0.34'],['风险标准差均值','0.4526','0.7624']])
r=record(33,'旧OT替换','旧 BLCA 两折均匀计划替换','两折 JSON 已保存 模型身份需核对','固定模型 monkey-patch OT 计划为均匀；不是训练时移除 OT。','两折 No-OT 分数略高；运行 method 名和 checkpoint 身份仍需核对，不外推到 v3.13 Full。',['paper_outputs/_archive_v310_v311/ot_contribution_blca_fold0.json','paper_outputs/_archive_v310_v311/ot_contribution_blca_fold1.json','paper_outputs/_archive_v310_v311/ot_contribution_summary.md'],[ARCH+'test_ot_contribution.py'])
vals=[j(x) for x in r['sources'][:2]];table(r,'两折计划替换',['Fold','Full C-index','均匀计划 C-index','Full 减均匀'],[[v['fold'],f"{v['full_cindex']:.4f}",f"{v['no_ot_cindex']:.4f}",f"{v['ot_contribution']:+.4f}"] for v in vals])
r=record(34,'旧统计检验','旧五癌种 DCT 和 SlotSPE 汇总统计','JSON 已保存 输入含已作废值','使用癌种汇总值作检验；非同患者配对外部复现。','旧 Wilcoxon p=0.625、t-test p=0.7284。输入 KIRC 0.8579 已失效，BLCA differences 字段也不等于两均值之差；不引用此检验为当前优劣证据。',['paper_outputs/_archive_v310_v311/statistical_test_results.json','STATISTICAL_ANALYSIS_COMPLETE_REPORT.md'],[ARCH+'statistical_significance_analysis.py'])
a=j(r['sources'][0]);table(r,'旧统计文件保留值',['项','值','状态'],[['Wilcoxon p',a['wilcoxon_p'],'历史输入 不作当前结论'],['t-test p',a['ttest_p'],'历史输入 不作当前结论'],['KIRC DCT 均值',a['dct_means']['kirc'],'两折旧值 作废'],['BLCA 差值字段',a['differences']['blca'],'与均值差不一致']])
r=record(35,'旧KM','旧版 KM 三组分层与五癌种曲线','汇总 CSV 已保存 原图部分本地缺失','旧版按风险三组或两组；不是 E011 的训练中位数十队列方案。','保留组数、人数、p 值和来源，旧版与新版不相互覆盖；原始患者预测、分组阈值与选择协议待补。',folder_sources('paper_outputs/_archive_v310_v311/km_curves')+folder_sources('paper_outputs/_archive_v310_v311/km_curves_multi_cancer'),[ARCH+'kaplan_meier_analysis.py',ARCH+'kp_multi_cancer.py'])
with locate('paper_outputs/_archive_v310_v311/km_curves_multi_cancer/km_summary.csv').open(encoding='utf-8-sig') as f:a=list(csv.reader(f));table(r,'旧五癌种 KM 摘要',a[0][:5],[v[:5] for v in a[1:]])
r=record(36,'隔离方向两折','六癌种 Direction OFF 和 ON 两折报告','明确隔离 不进入论文结果','SurvOT-Rank 271466d；legacy folds2、3；冻结类强制方向权重与 OFF 声明冲突。','原分数原样保留，仅用于追查历史，不认定为合法 OFF/ON 比较。',['experiments/candidate_evidence/DIRECTION_LOSS_2FOLD_LEGACY_VAL.md'],[ARCH+'run_dct_v310_experiments.py'])
parsed(r,r['sources'][0])
r=record(37,'撤回运输对照','v3.13 旧运输对照撤回档案','已撤回 禁止作为结果引用','2026-10-05 旧稿和比较图；折级重复问题与历史批次溯源未解决。','旧记录永久保留并标记撤回；不沿用旧图或旧控制臂数值。E008/E009 为后续重新审计的独立记录。',fig_sources('paper/history/retracted_transport_controls_20261005') if (FIGSRC/'paper/history').exists() else ['paper/history/retracted_transport_controls_20261005/README.md','paper/history/retracted_transport_controls_20261005/v313_transport_controls.png','paper/history/retracted_transport_controls_20261005/DCT_v313_初稿.md','paper/history/retracted_transport_controls_20261005/DCT_v313_初稿.docx'])
table(r,'撤回状态',['材料','处理','后续证据'],[['旧运输比较图及旧稿','保留 禁止引用','后续 controls_audit.json 单独登记'],['服务器旧结果目录','保留 未清空','未完成全量独立溯源']])
r=record(38,'v313多病例','v3.13 多病例解释图补录','已同步病例 PNG 缺部分逐病例身份清单','BLCA/KIRC 的 Full 与控制臂病例图；按目录和文件计数，文件数不是独立患者数。','上一版只展示一个病例，遗漏了两癌种多折病例图；已补归档。图片本身不证明真实组织空间对应。',fig_sources('paper/figures/case_interpretation')+fig_sources('paper/figures/case_interpretation_v2'),['scripts/plot_fig_c_case_interpretation.py'])
counts={}
for rel in r['sources']:
 if rel.endswith('.png'):
  key=str(Path(rel).parent).replace('\\','/');counts[key]=counts.get(key,0)+1
table(r,'病例图文件覆盖',['版本目录','病例 PNG 数','统计单位'],[[str(k).replace('paper/figures/',''),v,'图文件'] for k,v in counts.items()])
# Preserve additional original figures as material; do not duplicate derived plots as new training runs.
rels=['paper/figures/km_blca_pvalues.json','paper/figures/km_kirc_pvalues.json','paper/figures/fig3_risk_delta_heatmap.png','paper/figures/cohort_pathway_stats.json','paper/figures/cohort_pathway_heatmap_stats.json','paper/figures/km_blca_per_fold.png','paper/figures/km_kirc_per_fold.png','results/multi_cancer_summary_v313.json']+fig_sources('paper/figures/v313_main_plots')
r=record(39,'v313早期图','v3.13 逐折 KM 风险差值和早期图版本','有导出图与汇总 属于复用分析','记录被当前整合图替代的早期图、折级 KM 和风险差值热图。','已保留旧图版本与十癌种汇总 JSON；早期三折 sweep 汇总不替代 E010 的十折结果。早期解释矩阵需核对病例与权重定义。',rels,['scripts/plot_fig3_from_exports.py','scripts/plot_fig5_km_curves.py','scripts/prepare_v313_manuscript_panels.py'])
table(r,'补入材料',['类型','来源','使用状态'],[['逐折 KM','BLCA KIRC pvalues 与 per_fold PNG','旧版 留存'],['计划风险差值','fig3_risk_delta_heatmap.png','需核对数值来源'],['早期整合图','v313_main_plots','部分被整合版替代'],['十癌种均值 JSON','multi_cancer_summary_v313.json','E007 同批结果 不是新训练']])
r=record(40,'v314结构审计','v3.14 重建与槽结构数值审计','合成数据结构检查 真实癌种分数未建立','CPU/GPU、多 seed、masked/full/off 控制与维度修复；各 JSON 有自己的参数与源码哈希。','结构审计验证有限步前向、反向和有限值，不能声称 v3.14 的 C-index 优于 v3.13。',[p for p in folder_sources('audit_results') if '/v314_' in p]+['docs/DCT_V314_DEBUG_REPORT.md','docs/DCT_V314_MASKED_TRANSPORT_RECONSTRUCTION_PLAN.md'],[ARCH+'audit_v314.py'])
table(r,'合成检查覆盖',['材料','设备','模式','步骤'],[[Path(x).name,j(x).get('device','见配置'),j(x).get('mode','见配置'),j(x).get('steps','见配置')] for x in r['sources'] if x.endswith('.json')])
r=record(41,'v315结构与排序','v3.15 RTI 基线结构审计与排序实现','合成检查有 JSON 新排序实现无临床分数','旧 NLL-only 基线 CPU/GPU/independent 审计；RTI 定向排序是后续候选配置，证据分开。','旧 GPU 合成配置三 seed 各 30 步有限值；independent 的交互为零。后续 RTI 排序代码完成不等于有真实验证 C-index。',[p for p in folder_sources('audit_results') if '/v315_' in p]+['docs/DCT_V315_BASELINE.md','docs/DCT_V315_PLAN.md'],[ARCH+'audit_v315.py','configs/dct_v315_rti_rank_blca_uni2h.yaml'])
a=j('audit_results/v315_gpu.json');table(r,'GPU 合成基线摘要',['Seed','有限步数','NLL transport 梯度','峰值 allocated MiB'],[[k,v['finite_steps'],f"{v['nll_gradient_norms']['transport']:.4f}",f"{v['peak_allocated_mib']:.2f}"] for k,v in a['seeds'].items()])
r=record(42,'v316候选实现','v3.16 槽交互分解候选实现','本地代码和测试存在 未找到训练结果','配置与 standalone 实现保留；不把测试源文件存在当作测试已执行。','找到候选模块、配置和测试，未找到对应癌种折级训练曲线、checkpoint 或 score JSON。',[],['configs/dct_v316_slot_mi_decomposition.yaml','survot_rank/research/methods/dct_v316_slot_mi_interaction/model.py','survot_rank/research/methods/dct_v316_slot_mi_interaction/dct_v316_model.py','tests/test_dct_v316_slot_mi.py','tests/test_dct_v316_standalone.py'])
table(r,'当前证据',['材料','状态','可支持范围'],[['本地模型 配置 测试','存在','候选实现'],['真实训练曲线和分数','未找到','不填性能数值'],['历史测试完成声明','未在本次重跑','不替代运行记录']])
r=record(43,'运输修复演示','运输增强模块的短程演示','历史模块演示 无癌种验证结果','历史报告 3 秒 50 步模块优化，多分辨率锚点与对比目标。','报告时间变化指标 0.0059 至 0.3845，属于组件演示；不能写成真实生存预测提升或已集成训练成果。',['transport_fix_results.md','transport_fix_proposal.md'],[ARCH+'train_transport_fix_minimal.py',ARCH+'validate_transport_improvements.py'])
parsed(r,'transport_fix_results.md',0)
r=record(44,'待运行与缺失','待运行计划和未找到原始产物清单','计划或材料待补 不计完成训练','外层测试、外部泛化、正式匹配基线、risk ordering、双倍三倍重建与 reader 控制等。','计划必须保存，但未找到实际运行工件时不填分数、不计完成次数。服务器 /data1/DCT-Reg/results 本次未直接扫描。',['experiments/REGISTRY.csv','experiments/PROTOCOL.md','实验清单_最终版.md','还需要运行什么实验.md','还需要运行什么实验_简洁版.md','REMAINING_5FOLD_EXPERIMENTS.md','FINAL_EXPERIMENTS_TODO.md','COMPLETE_STATUS_SUMMARY.md'],['dct_v3_risk_ordering.py',ARCH+'ablate_reader.py',ARCH+'compute_calibration_and_dca.py',ARCH+'rerun_dct_v313_triple_w.sh',ARCH+'rerun_dct_v313_disable_both.sh'])
table(r,'尚不能登记为完成的方向',['方向','本次找到的材料','当前状态'],[['外层验证与外部泛化','REGISTRY 与 PROTOCOL','未找到完整证据'],['risk ordering','本地候选代码','未找到训练结果'],['重建权重变体 reader 控制','历史脚本','对应运行结果未同步'],['校准和决策曲线','历史代码与删除状态记录','原始指标未找到'],['实验服务器全量清单','/data1/DCT-Reg/results','未直接连接扫描']])
assert len(new)==28
extra=[x for x in folder_sources('audit_results') if '/v314_' not in x and '/v315_' not in x]
extra += [x for x in folder_sources('paper_outputs/_archive_v310_v311') if Path(x).name in ['run_summary.md','PAPER_COMPLETION_REPORT.md','table4_cross_cancer.csv','figure3_partial_data.json']]
for rel in extra:
 r=new[14];p=locate(rel);out=ASSETS/r['folder']/'原始材料'/rel;out.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,out)
 item=dict(experiment=r['id'],source=rel,source_kind='git_snapshot',source_commit=COMMIT,target=str(out.relative_to(BASE)),sha256=sha(out));r['sources'].append(rel);r['assets'].append(item);copied.append(item)

# One frozen supplementary code folder; per-experiment indexes resolve renamed archive scripts.
allpaths=set(subprocess.check_output(['git','ls-tree','-r','--name-only',COMMIT],cwd=ROOT).decode('utf8').splitlines())
snap=CODE/'补录代码快照_2b1c622';snap.mkdir(exist_ok=True);code_files=[]
for r in new:
 index=[]
 for rel in r['scripts']:
  if rel in allpaths:
   data=subprocess.check_output(['git','show',f'{COMMIT}:{rel}'],cwd=ROOT);out=snap/rel;kind='git_snapshot'
  elif (ROOT/rel).is_file():data=(ROOT/rel).read_bytes();out=CODE/r['folder']/'本地候选代码'/rel;kind='local_worktree'
  else:index.append(dict(requested=rel,present=False));continue
  out.parent.mkdir(parents=True,exist_ok=True);out.write_bytes(data)
  item=dict(requested=rel,present=True,path=str(out.relative_to(BASE)),sha256=sha(out),source_kind=kind);index.append(item);code_files.append(item)
 (CODE/r['folder']/'代码索引.json').write_text(json.dumps(dict(id=r['id'],entries=index,notice='归档时源代码与历史运行源码不一定相同。按原始产物内 source_sha256 和历史配置核对；本次未执行模型。'),ensure_ascii=False,indent=2),encoding='utf8')
 (CODE/r['folder']/'README.md').write_text('# '+r['id']+' '+r['title']+'\n\n代码索引.json 指向实际归档代码。present=false 表示未找到该入口，不伪造源文件。\n',encoding='utf8')
 for i,t in enumerate(r['tables'],1):
  with (ASSETS/r['folder']/f'表{i:02d}.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows([t['headers']]+t['rows'])
 (ASSETS/r['folder']/'实验材料索引.json').write_text(json.dumps(r,ensure_ascii=False,indent=2),encoding='utf8')
# Preserve all pre-existing root research reports, including uncertain and failed records.
reportdir=MAINT/'历史报告总库_工作区快照';reportdir.mkdir(exist_ok=True);reports=[]
for p in sorted(ROOT.glob('*.md')):
 if p.name in ['AGENTS.md','EXPERIMENTS.md']:continue
 out=reportdir/p.name;shutil.copy2(p,out);reports.append(dict(source=p.name,target=str(out.relative_to(BASE)),sha256=sha(out)))
(MAINT/'补录代码来源清单.json').write_text(json.dumps(code_files,ensure_ascii=False,indent=2),encoding='utf8')
(MAINT/'补录材料来源清单.json').write_text(json.dumps(copied+reports,ensure_ascii=False,indent=2),encoding='utf8')
(MAINT/'查漏范围.json').write_text(json.dumps(dict(source_commit=COMMIT,local_result_files_including_ignored=57,records_before=16,records_after=44,added_ids=[r['id'] for r in new],root_reports=len(reports),server_scanned=False,server_root='/data1/DCT-Reg/results',notice='条目数不是独立训练次数或患者数；历史、撤回、合成与待补记录分别标记。',preexisting_missing_tracked_files=subprocess.check_output(['git','ls-files','--deleted'],cwd=ROOT).decode('utf8').splitlines()),ensure_ascii=False,indent=2),encoding='utf8')
# Preserve the existing Word layout and append the new entries before the backlog.
doc=Document(DOC)
helper_tree=ast.parse((ROOT/'.review/experiment-ledger-20261009/build_archive.py').read_text(encoding='utf8'))
for node in helper_tree.body:
 if isinstance(node,ast.FunctionDef) and node.name in ['p','heading','link','native_table','picture']:
  exec(compile(ast.Module(body=[node],type_ignores=[]),'<existing_document_helpers>','exec'),globals())
anchor=next(p for p in doc.paragraphs if p.text=='待补项与修订记录')._p
previous=anchor.getprevious()
if previous is not None and previous.xpath('.//w:br[@w:type="page"]'):anchor=previous
for par in doc.paragraphs:
 if '本次建档共 16 个' in par.text:par.text=par.text.replace('本次建档共 16 个','当前建档共 44 个')
 if par.text.startswith('本文件是 DCT 实验事实的长期主档。'):par.text='本文件是 DCT 实验事实的长期主档。E001 至 E016 保留 v3.13 当前论文材料；E017 至 E044 补入历史实验、审计、撤回记录、候选结构与待补计划。不同版本和协议分别登记，以后在本文件继续追加。'
 if '材料按 E001 至 E016' in par.text:par.text=par.text.replace('E001 至 E016','E001 至 E044')
 if 'E017' in par.text:par.text=par.text.replace('E017','E045')
for t in doc.tables:
 for row in t.rows:
  if row.cells[0].text=='历史补录':row.cells[1].text='E017 至 E044 已补入历史与候选记录';row.cells[2].text='服务器全量目录仍待补'
start=list(doc.element.body)
heading('档案保护规则与本次查漏范围',new=True)
p('本档案、实验代码、配置、原始表格、图片、日志、预测、checkpoint、审计与历史稿件均属于受保护科研材料。未得到用户针对具体文件的明确授权，不得删除、移动、批量覆盖或以清理重复为由丢弃。失败、撤回、暂不入论文和旧版本也应保留。')
p('主档修改前保留版本副本；同一实验修订应写明原因，不能用新报告悄悄替换旧结果。材料冲突时以可核验原始产物和真实协议核对，不按总结文件名中的“最终”判定可信程度。')
p('本次核对本地结果目录，包括被忽略文件、历史报告及 origin/main 的 2b1c622 快照。查漏后共 44 个条目；条目是实验组或分析记录，不是 44 次完整训练。服务器未直接连接，不能保证尚未同步的训练结果无遗漏。')
p('查漏记录见 维护资料／查漏范围.json；原 33 页主档已保存到 维护资料／历史版本／20261009_补录前。旧报告总库保留所有现存根目录研究报告。')
for r in new:
 heading(r['id']+' '+r['title'],new=True)
 p('状态：'+r['status']);p('设置：'+r['setting'])
 for t in r['tables']:native_table(t)
 p('成果与判断：'+r['conclusion']);p('待核对：'+r['pending'])
 par=p('相关材料：');link(par,'表格与原始材料',str((ASSETS/r['folder']).relative_to(BASE)));par.add_run('    ');link(par,'代码索引',str((CODE/r['folder']/'代码索引.json').relative_to(BASE)))
 p('来源：'+'；'.join(r['sources'][:3])+('；其余见材料索引。' if len(r['sources'])>3 else '。'),'Caption')
# Move newly appended blocks together, without rewriting the previous experiment sections.
for node in list(doc.element.body):
 if node not in start and node.tag!=qn('w:sectPr'):anchor.addprevious(node)
# Add protection on the cover so readers see it immediately.
cover=next(p for p in doc.paragraphs if p.text.startswith('主档的解释必须服从'))._p
node=OxmlElement('w:p');run=OxmlElement('w:r');t=OxmlElement('w:t');t.text='保护要求：未经用户对具体材料的明确授权，不得删除、移动或批量覆盖档案及相关代码、结果和历史记录。';run.append(t);node.append(run);cover.addprevious(node)
# Add this revision to the existing native log table.
for t in doc.tables:
 if [c.text for c in t.rows[0].cells]==['日期','修订内容','依据']:
  row=t.add_row();row.cells[0].text='2026-10-09';row.cells[1].text='查漏补入 E017 至 E044 并建立材料保护规则';row.cells[2].text='本地原始产物 历史报告及 Git 2b1c622'
  for i,cell in enumerate(row.cells):
   original=cell._tc.get_or_add_tcPr();cell._tc.remove(original);cell._tc.insert(0,deepcopy(t.rows[1].cells[i]._tc.get_or_add_tcPr()))
   for par in cell.paragraphs:
    par.paragraph_format.space_after=Pt(0)
    for run in par.runs:run.font.size=Pt(10);run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
doc.save(DOC)
records.extend(new);(MAINT/'实验目录_建档快照.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf8')
protection='''# 实验资料保护与主档维护\n\n用户明确要求这些科研材料不能随意删除。本规则适用于实验档案、原始结果、实验代码与配置、图表、训练日志、预测、checkpoint、审计报告及历史版本。\n\n- 未取得用户针对具体材料的明确授权，不得删除、移动、批量覆盖，或以旧版本、重复、失败、撤回、暂不用于论文为由自动清理。\n- [DCT实验总档案.docx](实验档案/DCT实验总档案.docx) 为持续续写的实验事实主档，论文与汇报应先查此档案及原始产物。数值冲突按真实协议和可核验产物核对。\n- 修改主档前保留历史副本；已有实验编号不复用。新增实验使用新编号，分别保存代码索引和表格图片。\n- 撤回记录保留并标记，不恢复为有效结果；计划、合成结构检查与真实训练分别登记。\n- 本轮查漏是本地与已同步材料范围，不能把未连接的服务器目录称为已全量检查。\n'''
ag=ROOT/'AGENTS.md'
assert not ag.exists(),'Do not overwrite an existing instruction file.'
ag.write_text(protection,encoding='utf8')
(ROOT/'EXPERIMENTS.md').write_text('# DCT 实验事实主档入口\n\n长期主档为 [DCT实验总档案.docx](实验档案/DCT实验总档案.docx)。当前 44 个实验与分析条目，覆盖本地及已同步材料；不表示 44 次完整训练。\n\n未经用户对具体文件的明确授权，禁止删除、移动、批量覆盖档案及其代码、配置、原始结果、图表、日志、预测、checkpoint、审计与历史记录。旧版本、失败、撤回和暂不入论文材料均保留。规则亦写入 AGENTS.md。\n\n新增从 E045 续写，修改前保存历史副本，再更新 Word 目录。原 16 条 v3.13 记录保留。服务器 /data1/DCT-Reg/results 尚未直接全量扫描。\n',encoding='utf8')
(BASE/'保护与续写规则.md').write_text(protection,encoding='utf8')
with (BASE/'README.md').open('a',encoding='utf8') as f:f.write('\n\n2026-10-09 查漏更新：44 条记录，新增 E017–E044；新实验从 E045 继续。所有档案和相关材料受保护，规则见 保护与续写规则.md 和根目录 AGENTS.md。旧主档保留在 维护资料/历史版本/。\n')
for rel,h in protected.items():
 if rel=='EXPERIMENTS.md':continue
 assert (ROOT/rel).is_file() and sha(ROOT/rel)==h,rel
print(json.dumps(dict(records=44,added=28,archived_materials=len(copied),root_reports=len(reports),code_entries=len(code_files),protected_existing_files_checked=len(protected),native_tables=len(doc.tables),images=len(doc.inline_shapes)),ensure_ascii=False))