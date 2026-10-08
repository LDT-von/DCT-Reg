#!/usr/bin/env python3
"""Render the v3.13 ten-cancer comparison table from saved, audited scores.

No training/inference. Complete input is required by default. An explicitly
requested draft shows missing cells and never invents winners for incomplete
columns. Matched cells are five-fold means; raw folds/std remain in the JSON.
--reported explicitly renders published means as a cross-protocol reference table.
"""
from __future__ import annotations
import argparse
import hashlib
import html
import json
import math
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from statistics import stdev

ROOT = Path(__file__).resolve().parents[1]
CANCERS = ("BRCA", "COADREAD", "KIRC", "UCEC", "LUAD",
           "LUSC", "HNSC", "SKCM", "BLCA", "STAD")
RED = "#D7191C"
YELLOW = "#FFF7CF"


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def quantize(value, digits):
    return Decimal(str(value)).quantize(Decimal(1).scaleb(-digits),
                                       rounding=ROUND_HALF_UP)


def dense_ranks(values, digits=4):
    """Tie equal displayed values; second distinct value receives rank 2."""
    shown = [quantize(v, digits) if v is not None else None for v in values]
    levels = sorted(set(v for v in shown if v is not None), reverse=True)
    return [levels.index(v) + 1 if v is not None else None for v in shown]


def summarize(payload, *, base=ROOT, digits=4):
    if payload.get("schema_version") != 1:
        raise ValueError("Expected schema_version=1")
    if tuple(payload.get("cancers", [])) != CANCERS:
        raise ValueError("All ten cancer columns must follow the declared table order")
    models = payload.get("models", [])
    ids = [m["id"] for m in models]
    if len(models) < 2 or len(set(ids)) != len(ids) or "dct_v313" not in ids:
        raise ValueError("Require unique model IDs, DCT v3.13 and at least one baseline")
    if digits not in (3, 4):
        raise ValueError("Display precision must be 3 or 4")
    protocol = payload.get("protocol", {})
    protocol_ready = protocol.get("audit") == "verified"
    protocol_errors = []
    for cancer in CANCERS:
        hashes = protocol.get("split_sha256", {}).get(cancer)
        if not isinstance(hashes, list) or len(hashes) != 5 or any(
                not isinstance(h, str) or len(h) != 64 or
                any(ch not in "0123456789abcdef" for ch in h) for h in hashes):
            protocol_errors.append(f"{cancer} split hashes")
    if protocol_errors:
        protocol_ready = False
    rows, problems = [], []
    if not protocol_ready:
        problems.append("Common protocol/split audit pending")
    for model in models:
        cells, ready = {}, {}
        row_ready = model.get("audit") == "verified" and model.get("protocol") == protocol
        if not row_ready:
            problems.append(f"{model['model']}: protocol/source audit pending")
        missing = []
        for cancer in CANCERS:
            folds = model.get("folds", {}).get(cancer)
            if folds is None:
                cells[cancer], ready[cancer] = None, False
                missing.append(cancer)
                continue
            if (not isinstance(folds, list) or len(folds) != 5 or
                    any(isinstance(v, bool) or not isinstance(v, (float, int)) or
                        not math.isfinite(v) or not 0 <= v <= 1 for v in folds)):
                raise ValueError(f"{model['model']} {cancer}: five finite C-index values in [0,1] required")
            mean = sum(Decimal(str(v)) for v in folds) / 5
            cells[cancer] = {"mean": float(mean), "std": stdev(folds),
                             "folds": folds, "text": str(quantize(mean, digits)),
                             "rank": None}
            sources = model.get("sources", {}).get(cancer, [])
            source_ready = (row_ready and protocol_ready and len(sources) == 5 and
                            sorted(s.get("fold", -1) for s in sources) == list(range(5)))
            if source_ready:
                for record in sources:
                    path = Path(record.get("path", ""))
                    if not path.is_absolute():
                        path = base / path
                    if (record.get("audit") != "verified" or not path.is_file() or
                            record.get("sha256") != sha256(path)):
                        source_ready = False
                        break
            ready[cancer] = source_ready
            if not source_ready and row_ready:
                problems.append(f"{model['model']} {cancer}: five verified source records required")
        if missing:
            problems.append(f"{model['model']}: missing {', '.join(missing)}")
        means = [cells[c]["mean"] for c in CANCERS if cells[c] is not None]
        if len(means) == 10:
            # Average ten unrounded cohort means, never the displayed numbers.
            overall = sum(sum(Decimal(str(v)) for v in cells[c]["folds"]) / 5
                          for c in CANCERS) / 10
            cells["Overall"] = {"mean": float(overall),
                                "text": str(quantize(overall, digits)), "rank": None}
        else:
            cells["Overall"] = None
        ready["Overall"] = all(ready.values())
        rows.append({"id": model["id"], "model": model["model"],
                     "modality": model["modality"], "group": model["group"],
                     "cells": cells, "ready": ready})
    ranked_columns = []
    for column in (*CANCERS, "Overall"):
        if all(row["ready"][column] and row["cells"][column] is not None for row in rows):
            ranks = dense_ranks([row["cells"][column]["mean"] for row in rows], digits)
            for row, rank in zip(rows, ranks):
                row["cells"][column]["rank"] = rank
            ranked_columns.append(column)
    ours = next(row for row in rows if row["id"] == "dct_v313")
    return {"complete": len(ranked_columns) == 11, "digits": digits,
            "cancers": list(CANCERS), "rows": rows, "problems": problems,
            "ranked_columns": ranked_columns,
            "ranking": "dense ranks of displayed rounded means; ties share rank",
            "overall": "equal mean of ten unrounded cohort means",
            "dct_first": [c for c in ranked_columns if ours["cells"][c]["rank"] == 1],
            "dct_second": [c for c in ranked_columns if ours["cells"][c]["rank"] == 2]}


def summarize_reported(payload, *, digits=3):
    """Rank source-reported means without inventing patient records or folds."""
    if payload.get("comparison_type") != "reported_reference":
        raise ValueError("--reported requires comparison_type=reported_reference")
    if payload.get("schema_version") != 1 or digits not in (3, 4):
        raise ValueError("Require schema_version=1 and display precision 3 or 4")
    if tuple(payload.get("cancers", [])) != CANCERS or payload.get("transcription_verified") is not True:
        raise ValueError("All ten source columns and verified transcription required")
    models = payload["models"]
    if len(models) < 2 or len({m["id"] for m in models}) != len(models) or "dct_v313" not in {m["id"] for m in models}:
        raise ValueError("Require at least two unique model IDs")
    rows = []
    for model in models:
        cells = {}
        if model.get("source_kind") not in ("published_summary", "author_fold_summary"):
            raise ValueError("Explicit published/author source kind required")
        for column in (*CANCERS, "Overall"):
            value = model["reported_overall"] if column == "Overall" else model["means"][column]
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or not 0 <= value <= 1:
                raise ValueError(f"{model['model']} {column}: finite reported C-index required")
            cells[column] = {"mean": value, "text": str(quantize(value, digits)), "rank": None,
                             "std": model.get("std", {}).get(column),
                             "folds": model.get("folds", {}).get(column),
                             "source_kind": model["source_kind"]}
        rows.append({"id": model["id"], "model": model["model"],
                     "modality": model["modality"], "group": model["group"],
                     "cells": cells, "ready": {c:True for c in (*CANCERS, "Overall")}})
    for column in (*CANCERS, "Overall"):
        ranks = dense_ranks([r["cells"][column]["mean"] for r in rows], digits)
        for row, rank in zip(rows, ranks):
            row["cells"][column]["rank"] = rank
    ours = next(r for r in rows if r["id"] == "dct_v313")
    return {"complete":True, "matched_complete":False, "comparison_type":"reported_reference",
            "digits":digits, "cancers":list(CANCERS), "rows":rows, "problems":[],
            "ranked_columns":list((*CANCERS, "Overall")),
            "ranking":"dense ranks of displayed reported means; reference ranking only",
            "overall":"Baselines: original paper Overall; DCT: equal mean of unrounded cohort means",
            "source":payload["source"], "protocols":payload["protocols"],
            "reference_sample_sizes":payload["reference_sample_sizes"],
            "dct_first":[c for c in (*CANCERS, "Overall") if ours["cells"][c]["rank"] == 1],
            "dct_second":[c for c in (*CANCERS, "Overall") if ours["cells"][c]["rank"] == 2]}


def cell_text(cell, missing="--"):
    return cell["text"] if cell else missing


def notes(report):
    if report.get("notes"):
        return report["notes"]
    if report.get("comparison_type") == "reported_reference":
        return [
          "REFERENCE: baselines from SlotSPE Table 1 (UNI); DCT from author five-fold records (UNI2-h).",
          "Encoder and patient splits are not aligned; ranks compare reported values, not a matched reproduction.",
          "Best: red bold. Second: underlined. Ties share rank. Overall: paper values for baselines; unrounded macro mean for DCT.",
          "g. = genomics; h. = histology. Paper cohort counts are not assigned to the DCT row. Source means/SD are retained in JSON."]
    status = ("Complete source manifest: shared protocol declared and source-file hashes checked."
              if report["complete"] else
              "DRAFT: baseline results/source audits pending; missing cells are --; incomplete columns are not ranked.")
    return [status,
            "Cells: five-fold mean C-index. Overall: equal mean of ten unrounded cohort means. Fold values and SD: source JSON.",
            "Best: red bold. Second: underlined. Ties at displayed precision share a dense rank. Yellow row: DCT v3.13.",
            "g. = genomics; h. = histology; g.+h. = both. Patient counts are added only after common cohort identity is audited."]


def render_vector(report, stem):
    columns = tuple(report.get("columns", (*CANCERS, "Overall")))
    title = report.get("title", "DCT v3.13 Ten Cancer Model Comparison")
    from reportlab.pdfgen import canvas
    from reportlab.lib.colors import HexColor
    from reportlab.pdfbase.pdfmetrics import stringWidth
    width, left, right, top, row_height = 1400, 35, 1365, 126, 29
    model_width = report.get("model_column_width", 160)
    col_widths = [model_width, 94] + [(right-left-model_width-94)/len(columns)] * len(columns)
    edges = [left]
    for value in col_widths:
        edges.append(edges[-1] + value)
    height = top + row_height * len(report["rows"]) + max(128, 30 + 22*len(notes(report)))
    cv = canvas.Canvas(str(stem.with_suffix(".pdf")), pagesize=(width*.6, height*.6))
    cv.scale(.6, .6)
    svg = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
           '<rect width="100%" height="100%" fill="white"/>']
    def line(x1, y1, x2, y2, weight=1, color="#111111"):
        svg.append(f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="{color}" stroke-width="{weight}"/>')
        cv.setStrokeColor(HexColor(color)); cv.setLineWidth(weight)
        cv.line(x1, height-y1, x2, height-y2)
    def text(x, y, value, size=18, color="#111111", bold=False, underline=False):
        family = "Times-Bold" if bold else "Times-Roman"
        cv.setFillColor(HexColor(color)); cv.setFont(family, size)
        cv.drawCentredString(x, height-y, value)
        svg.append(f'<text x="{x}" y="{y}" text-anchor="middle" font-family="Times New Roman,serif" font-size="{size}" font-weight="{"bold" if bold else "normal"}" fill="{color}">{html.escape(value)}</text>')
        if underline:
            span = stringWidth(value, family, size)
            line(x-span/2, y+3, x+span/2, y+3, .9, color)
    text(width/2, 33, title, 25, bold=True)
    text(width/2, 62, "Best", 17, RED, True)
    text(width/2+62, 62, "Second", 17, underline=True)
    line(left, 82, right, 82, 2.2); line(left, 86, right, 86, .8)
    for j, header in enumerate((report.get("model_header", "Model"), report.get("input_header", "Modality"), *columns)):
        text((edges[j]+edges[j+1])/2, 111, header, 17.2, bold=True)
    line(left, top, right, top, 1.6)
    previous_group = None
    for i, row in enumerate(report["rows"]):
        y = top + row_height*i
        if row["id"] == "dct_v313":
            svg.append(f'<rect x="{left}" y="{y}" width="{right-left}" height="{row_height}" fill="{YELLOW}"/>')
            cv.setFillColor(HexColor(YELLOW)); cv.rect(left, height-y-row_height, right-left, row_height, fill=1, stroke=0)
            line(left, y, right, y, 1.5)
        elif previous_group is not None and row["group"] != previous_group:
            line(left, y, right, y, .7, "#555555")
        previous_group = row["group"]
        ours = row["id"] == "dct_v313"
        text((edges[0]+edges[1])/2, y+21, row["model"], 18, bold=ours)
        text((edges[1]+edges[2])/2, y+21, row["modality"], 18)
        for j, column in enumerate(columns, 2):
            cell = row["cells"][column]; rank = cell["rank"] if cell else None
            text((edges[j]+edges[j+1])/2, y+21, cell_text(cell, report.get("missing_label", "--")), 18,
                 RED if rank == 1 else "#111111", ours or rank == 1, rank == 2)
    bottom = top + row_height*len(report["rows"])
    line(left, bottom, right, bottom, 2)
    for i, value in enumerate(notes(report)):
        # Footnotes are left-aligned to keep the visual compact.
        cv.setFont("Times-Roman", 15); cv.setFillColor(HexColor("#333333"))
        cv.drawString(left, height-bottom-25-i*22, value)
        svg.append(f'<text x="{left}" y="{bottom+25+i*22}" font-family="Times New Roman,serif" font-size="15" fill="#333333">{html.escape(value)}</text>')
    svg.append("</svg>");stem.with_suffix(".svg").write_text("\n".join(svg), encoding="utf-8",newline="\n")
    cv.showPage(); cv.save()
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(stem.with_suffix(".pdf")) as doc:
        page=doc[0]; bitmap=page.render(scale=2.8)
        bitmap.to_pil().save(stem.with_suffix(".png"));bitmap.close();page.close()


def render_docx(report, stem):
    columns = tuple(report.get("columns", (*CANCERS, "Overall")))
    title_text = report.get("title", "DCT v3.13 Ten Cancer Model Comparison")
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    doc = Document(); sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21)
    sec.left_margin = sec.right_margin = Cm(1.2)
    sec.top_margin = sec.bottom_margin = Cm(1.3)
    for style in ("Normal", "Title"):
        st=doc.styles[style];st.font.name="Times New Roman"
        st.font.color.rgb=RGBColor(0,0,0);st.font.underline=False
        st.font.size=Pt(9 if style=="Normal" else 16)
        st.paragraph_format.space_after=Pt(4)
    title=doc.add_paragraph(title_text, style="Title")
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in title.runs:
        r.font.name="Times New Roman";r.font.bold=True
        r._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),"Times New Roman")
        r._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),"Times New Roman")
    # Remove inherited title-rule residue from the bundled default template.
    for parent in (doc.styles._element,doc._element):
        for border in list(parent.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run("Best");r.bold=True;r.font.color.rgb=RGBColor.from_string(RED[1:])
    p.add_run("     ");p.add_run("Second").underline=True
    table=doc.add_table(rows=0,cols=2+len(columns));table.autofit=False
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    model_cm=report.get("model_column_cm",3.5)
    widths=[model_cm,2.0]+[(27.3-model_cm-2.0)/len(columns)]*len(columns)
    for col,w in zip(table.columns,widths):col.width=Cm(w)
    data=[[report.get("model_header", "Model"),report.get("input_header", "Modality"),*columns]]+[
        [row["model"],row["modality"],*[cell_text(row["cells"][c], report.get("missing_label", "--")) for c in columns]]
        for row in report["rows"]]
    last_group=None
    for i, values in enumerate(data):
        row=table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        info=report["rows"][i-1] if i else None
        ours=bool(info and info["id"]=="dct_v313")
        separator=bool(info and last_group is not None and info["group"]!=last_group)
        if info:last_group=info["group"]
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement("w:tcBorders")
            for edge in ("top","bottom","left","right","insideH","insideV"):
                node=OxmlElement("w:"+edge);node.set(qn("w:val"),"nil")
                if (edge=="top" and (i==0 or ours or separator)) or (edge=="bottom" and (i==0 or i==len(data)-1)):
                    node.set(qn("w:val"),"single");node.set(qn("w:sz"),"10" if ours or i==0 else "5")
                    node.set(qn("w:color"),"111111")
                borders.append(node)
            tcpr.append(borders)
            margins=OxmlElement("w:tcMar")
            for side in ("top","bottom","left","right"):
                item=OxmlElement("w:"+side);item.set(qn("w:w"),"30" if report.get("compact") else "45");item.set(qn("w:type"),"dxa");margins.append(item)
            tcpr.append(margins)
            if ours:
                sh=OxmlElement("w:shd");sh.set(qn("w:fill"),YELLOW[1:]);tcpr.append(sh)
            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pf=p.paragraph_format;pf.space_before=pf.space_after=Pt(1 if report.get("compact") else 2)
            pf.line_spacing=1.0;pf.keep_with_next=i!=len(data)-1
            r=p.add_run(value);r.font.name="Times New Roman"
            r.font.size=Pt(report.get("header_font_sizes",{}).get(value,9.2) if i==0 else 9.2)
            rank=info["cells"][columns[j-2]]["rank"] if info and j>=2 and info["cells"][columns[j-2]] else None
            r.bold=i==0 or ours or rank==1;r.underline=rank==2
            r.font.color.rgb=RGBColor.from_string(RED[1:] if rank==1 else "111111")
    for value in notes(report):
        p=doc.add_paragraph(value);p.paragraph_format.space_after=Pt(3)
        for r in p.runs:r.font.size=Pt(8.5)
    doc.save(stem.with_suffix(".docx"))


def render_html(report, stem):
    columns = tuple(report.get("columns", (*CANCERS, "Overall")))
    title = report.get("title", "DCT v3.13 Ten Cancer Model Comparison")
    rows=[]
    for row in report["rows"]:
        cells=[f"<td>{html.escape(row['model'])}</td>",f"<td>{html.escape(row['modality'])}</td>"]
        for column in columns:
            cell=row["cells"][column];rank=cell["rank"] if cell else None
            cls="best" if rank==1 else "second" if rank==2 else ""
            cells.append(f'<td class="{cls}">{cell_text(cell, report.get("missing_label", "--"))}</td>')
        rows.append(f'<tr class="{"ours" if row["id"]=="dct_v313" else ""}">{"".join(cells)}</tr>')
    headers="".join(f"<th>{c}</th>" for c in (report.get("model_header", "Model"),report.get("input_header", "Modality"),*columns))
    foot="".join(f"<p>{html.escape(n)}</p>" for n in notes(report))
    content=f'''<!doctype html><html><meta charset="utf-8"><title>DCT v3.13 comparison</title>
<style>body{{font-family:"Times New Roman",serif;padding:24px}}h1{{font-size:24px}}table{{border-collapse:collapse;width:100%;min-width:1240px;border-top:3px double;border-bottom:2px solid}}th,td{{text-align:center;padding:7px 9px;font-size:16px}}thead{{border-bottom:1.5px solid}}.best{{color:{RED};font-weight:bold}}.second{{text-decoration:underline}}.ours{{background:{YELLOW};font-weight:bold;border-top:1.5px solid}}p{{font-size:13px}}@media print{{@page{{size:landscape;margin:12mm}}body{{padding:0}}table{{min-width:0}}}}</style>
<h1>{html.escape(title)}</h1><p><b style="color:{RED}">Best</b> &nbsp; <u>Second</u></p>
<table><thead><tr>{headers}</tr></thead><tbody>{"".join(rows)}</tbody></table>{foot}</html>'''
    stem.with_suffix(".html").write_text(content,encoding="utf-8",newline="\n")


def main(argv=None):
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input",type=Path,default=ROOT/"paper/V313_TEN_CANCER_COMPARISON_INPUT.json")
    p.add_argument("--output",type=Path,default=ROOT/"paper/tables/v313_ten_cancer_comparison")
    p.add_argument("--digits",type=int,choices=(3,4),default=4)
    p.add_argument("--allow-incomplete",action="store_true",help="Explicit draft preview; never ranks incomplete columns")
    p.add_argument("--check-only",action="store_true")
    p.add_argument("--reported",action="store_true",help="Explicit cross-publication reported-value reference; never a matched comparison")
    args=p.parse_args(argv)
    try:
        payload=json.loads(args.input.read_text(encoding="utf-8"))
        report=(summarize_reported(payload,digits=args.digits) if args.reported
                else summarize(payload,digits=args.digits))
        if not report["complete"] and not args.allow_incomplete:
            raise ValueError("Comparison incomplete; no files written:\n"+"\n".join(report["problems"]))
        if args.check_only:
            print(json.dumps({"complete":report["complete"],"ranked_columns":report["ranked_columns"],"problems":report["problems"]},ensure_ascii=False,indent=2));return 0
        args.output.mkdir(parents=True,exist_ok=True)
        stem=args.output/"table_v313_ten_cancer_comparison"
        render_vector(report,stem);render_docx(report,stem);render_html(report,stem)
        report["input_sha256"]=sha256(args.input)
        report["outputs"]={suffix:sha256(stem.with_suffix("."+suffix)) for suffix in ("pdf","svg","png","docx","html")}
        stem.with_suffix(".json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8",newline="\n")
        print(f"Table: {stem}; complete={report['complete']}; ranked columns={len(report['ranked_columns'])}/11")
    except (ValueError,OSError,KeyError,ImportError) as error:
        p.exit(1,f"{error}\nNo model was run.\n")
    return 0


if __name__=="__main__":
    raise SystemExit(main())
