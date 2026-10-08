"""Build the v3.13 manuscript and figures from recorded experiment values.

Uses the Codex bundled Python runtime; never starts model training.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

import numpy as np
import pypdfium2 as pdfium
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.opc.constants import RELATIONSHIP_TYPE
from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing, Rect, String, Line, Polygon, Circle
from reportlab.lib.colors import HexColor
from draw_v313_architecture import build_architecture

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'paper' / 'DCT_v313_初稿.md'
OUTPUT = SOURCE.with_suffix('.docx')
FIGURES = ROOT / 'paper' / 'figures'
QA = ROOT / 'tmp' / 'v313_paper_qa'
BLCA = np.array([
    [.6423,.6721,.6988,.7103,.7769], [.6296,.6918,.7044,.7340,.7889],
    [.6517,.6927,.7147,.7954,.7402], [.6576,.7030,.7399,.7524,.7496],
    [.6576,.7056,.6810,.7629,.7436], [.6576,.6910,.7203,.7805,.7444],
    [.6576,.6953,.7465,.7744,.7453],
])
KIRC = np.array([
    [.8073,.8640,.8379,.7837,.7671], [.7886,.8471,.7834,.7958,.7524],
    [.7800,.8268,.7752,.7943,.7841], [.8203,.8703,.7698,.7905,.7738],
    [.8239,.8527,.7732,.7890,.7841], [.7951,.8485,.7738,.7927,.8032],
    [.8253,.8485,.8202,.8094,.8084],
])
NAVY, TEAL, RED = '#244b70', '#007f83', '#ad4937'


def label(d, x, y, text, size=9, color='#202b36', anchor='start', bold=False):
    d.add(String(x, y, text, fontName='Helvetica-Bold' if bold else 'Helvetica',
                 fontSize=size, fillColor=HexColor(color), textAnchor=anchor))


def arrow(d, x1, y1, x2, y2, color='#65778a', dashed=False):
    ln = Line(x1, y1, x2, y2, strokeColor=HexColor(color), strokeWidth=1.2)
    if dashed:
        ln.strokeDashArray = [4, 3]
    d.add(ln)
    if abs(x2-x1) >= abs(y2-y1):
        sign = 1 if x2 > x1 else -1
        points = [x2,y2,x2-sign*6,y2+3,x2-sign*6,y2-3]
    else:
        sign = 1 if y2 > y1 else -1
        points = [x2,y2,x2-3,y2-sign*6,x2+3,y2-sign*6]
    d.add(Polygon(points, fillColor=HexColor(color), strokeColor=None))


def box(d, x, y, w, h, rows, fill='#f0f4f7', color=NAVY):
    d.add(Rect(x,y,w,h,rx=5,ry=5,fillColor=HexColor(fill),
               strokeColor=HexColor(color),strokeWidth=.8))
    for i,text in enumerate(rows):
        label(d,x+w/2,y+h/2+(len(rows)-1)*6-i*12,text,8.5,color,'middle',i==0)


def save_drawing(d, name):
    pdf = QA / (name+'.pdf')
    renderPDF.drawToFile(d,str(pdf))
    doc = pdfium.PdfDocument(str(pdf))
    doc[0].render(scale=3).to_pil().save(FIGURES/(name+'.png'))
    doc.close()


def make_figures():
    FIGURES.mkdir(parents=True,exist_ok=True)
    QA.mkdir(parents=True,exist_ok=True)
    build_architecture()

    d=Drawing(540,248)
    label(d,12,232,'Loss recipe comparison   Five-fold mean best-validation C-index',11,bold=True)
    for x,arr,title,low,high,color in [(45,BLCA,'BLCA',.69,.735,NAVY),(311,KIRC,'KIRC',.785,.83,TEAL)]:
        width,height,y=207,151,48
        label(d,x+width/2,210,title,11,color,'middle',True)
        for value in np.linspace(low,high,4):
            py=y+(value-low)/(high-low)*height
            d.add(Line(x,py,x+width,py,strokeColor=HexColor('#dbe1e7'),strokeWidth=.6))
            label(d,x-6,py-3,f'{value:.3f}',8,anchor='end')
        d.add(Line(x,y,x,y+height,strokeColor=HexColor('#4b5964')))
        d.add(Line(x,y,x+width,y,strokeColor=HexColor('#4b5964')))
        for j,value in enumerate(arr.mean(axis=1)):
            px=x+12+j*(width-24)/6
            py=y+(value-low)/(high-low)*height
            c=RED if j==6 else color
            d.add(Circle(px,py,4,fillColor=HexColor(c),strokeColor=None))
            label(d,px,y-15,'Full' if j==6 else f'E{j}',8,anchor='middle')
            label(d,px,py+8,f'{value:.4f}',7.4,c,'middle')
    label(d,270,14,'E4 self and E5 cross both branch from E3; they are not successive additions.',8,anchor='middle')
    save_drawing(d,'v313_loss_ablation')

# Small deterministic TeX -> OMML translator for this manuscript's equations.
# Fractions, scripts and accents are native editable Word math nodes.
SYMBOLS={'prod':'∏','sum':'∑','leq':'≤','geq':'≥','langle':'⟨','rangle':'⟩',
         'tau':'τ','epsilon':'ε','varepsilon':'ε','top':'⊤','odot':'⊙','sigma':'σ','rho':'ρ',
         'quad':'  ','qquad':'    ','min':'min','arg':'arg','|':'‖',
         'left':'','right':'','!':'',',':' ', ';':' ', 'colon':':',
         'in':'∈','times':'×','delta':'δ','alpha':'α','mathsf':'','cdot':'·','{':'{','}':'}'}


def math_run(text, plain=False):
    r=OxmlElement('m:r')
    if plain:
        pr=OxmlElement('m:rPr'); st=OxmlElement('m:sty'); st.set(qn('m:val'),'p'); pr.append(st);r.append(pr)
    pr=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for attr in ('ascii','hAnsi'):
        fonts.set(qn('w:'+attr),'Cambria Math')
    pr.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);r.append(pr)
    t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=text;r.append(t)
    return r


def math_element(tag,nodes):
    e=OxmlElement(tag)
    for node in nodes: e.append(node)
    return e


class MathParser:
    def __init__(self,text): self.text=text;self.i=0
    def group(self):
        if self.i<len(self.text) and self.text[self.i]=='{':
            self.i+=1; nodes=self.sequence('}');self.i+=1;return nodes
        return [self.atom()]
    def atom(self):
        c=self.text[self.i]
        if c=='{':
            return math_element('m:box',[math_element('m:e',self.group())])
        if c!='\\':
            self.i+=1;return math_run(c)
        self.i+=1
        m=re.match(r'[A-Za-z]+|.',self.text[self.i:]); cmd=m.group();self.i+=len(cmd)
        if cmd in ('frac','tfrac'):
            numerator=self.group();denominator=self.group()
            return math_element('m:f',[math_element('m:num',numerator),math_element('m:den',denominator)])
        if cmd in ('bar','widetilde','widehat'):
            pr=OxmlElement('m:accPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),{'bar':'̅','widetilde':'̃','widehat':'̂'}[cmd]);pr.append(ch)
            return math_element('m:acc',[pr,math_element('m:e',self.group())])
        if cmd in ('operatorname','mathrm','mathcal','mathbf'):
            nodes=self.group()
            if cmd in ('operatorname','mathrm'):
                return math_run(''.join(n.text or '' for node in nodes for n in node.iter(qn('m:t'))),True)
            return math_element('m:box',[math_element('m:e',nodes)])
        if cmd not in SYMBOLS:
            raise ValueError('Unsupported TeX command '+cmd)
        return math_run(SYMBOLS[cmd],cmd in ('min','arg'))
    def sequence(self,stop=''):
        nodes=[]
        while self.i<len(self.text) and (not stop or self.text[self.i]!=stop):
            c=self.text[self.i]
            if c in '^_':
                scripts={}
                while self.i<len(self.text) and self.text[self.i] in '^_':
                    kind=self.text[self.i];self.i+=1;scripts[kind]=self.group()
                base=nodes.pop()
                tag='m:sSubSup' if len(scripts)==2 else ('m:sSub' if '_' in scripts else 'm:sSup')
                node=OxmlElement(tag);node.append(math_element('m:e',[base]))
                for kind,t in (('_','m:sub'),('^','m:sup')):
                    if kind in scripts:node.append(math_element(t,scripts[kind]))
                nodes.append(node)
            else:nodes.append(self.atom())
        return nodes


def native_math(tex):
    root=OxmlElement('m:oMath')
    for node in MathParser(tex).sequence():root.append(node)
    return root


def font(run,size=10.5,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold
    run.font.color.rgb=__import__('docx.shared',fromlist=['RGBColor']).RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')


def add_inline(p,text,size=10.5,bold=False):
    pieces=re.split(r'(\$[^$]+\$|\*\*.*?\*\*|\[[^\]]+\]\(https?://[^)]+\))',text)
    for s in pieces:
        if not s:continue
        if s.startswith('$') and s.endswith('$'):p._p.append(native_math(s[1:-1]))
        elif s.startswith('**'):add_inline(p,s[2:-2],size,True)
        elif re.match(r'\[[^\]]+\]\(https?://',s):
            label_,url=re.match(r'\[([^\]]+)\]\(([^)]+)\)',s).groups()
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RELATIONSHIP_TYPE.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label_;r.append(t);link.append(r);p._p.append(link)
        else:font(p.add_run(s),size,bold)


def add_equation(doc,tex):
    number=re.search(r'\\tag\{(\d+)\}',tex).group(1)
    tex=re.sub(r'\\tag\{\d+\}','',tex).strip()
    # Explicit balanced lines keep longer equations within the text width.
    breaks={
        '2':r',\qquad s', '3':r',\qquad C', '4':r',\qquad T',
        '6':r',\qquad h', '7':r',\qquad \widetilde s^w_l',
        '9':r'+\operatorname{SmoothL1}',
        '11':r'+0.10r(e)',
        '13':r',\qquad T_\alpha', '14':r',\qquad H_l',
    }
    split=breaks.get(number)
    lines=tex.split(split,1) if split and split in tex else [tex]
    if len(lines)==2:
        if split.startswith(','):
            lines[0]+=',';lines[1]=split.removeprefix(r',\qquad ')+lines[1]
        else:lines[1]=split+lines[1]
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.space_before=Pt(3);pf.space_after=Pt(6);pf.keep_together=True
    if len(lines)==1:
        p._p.append(native_math(lines[0]))
    else:
        math=OxmlElement('m:oMath');arr=OxmlElement('m:eqArr')
        for line in lines:arr.append(math_element('m:e',MathParser(line).sequence()))
        math.append(arr);p._p.append(math)
    font(p.add_run('  ('+number+')'),9)


def add_table(doc,rows):
    n=len(rows[0]);widths={3:[1.1,3.0,12.6],4:[3.4,3.3,5.0,5.0],
        5:[3.25,3.35,3.35,3.2,3.55],6:[3.5,2.64,2.64,2.64,2.64,2.64],
        7:[2.8,2.0,2.0,2.0,2.0,2.0,3.9]}[n]
    if rows[0][0]=='Fold':widths=[1.4,2.5,2.5,3.2,3.35,3.75]
    t=doc.add_table(rows=0,cols=n);t.autofit=False;t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    keep_table_together = rows[0][:2] in (['配置', 'BLCA C-index'], ['配置', 'IPCW 排序'])
    for i,values in enumerate(rows):
        row=t.add_row();props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        if i==0:props.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=row.cells[j];cell.width=Cm(widths[j]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            pr=cell._tc.get_or_add_tcPr()
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EDEFF2');pr.append(sh)
            borders=OxmlElement('w:tcBorders')
            for side in ('top','bottom','left','right'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            pr.append(borders)
            margin=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margin.append(el)
            pr.append(margin)
            p=cell.paragraphs[0];pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.line_spacing=1.08;pf.space_after=Pt(4);pf.space_before=Pt(4)
            pf.keep_with_next=keep_table_together or i==0 or i==len(rows)-1 or len(rows)<=4
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p,value,8.5 if n==7 else (9.0 if n==6 else 9.5),i==0)


def add_figure_placeholder(doc, number):
    """Reserve an editable blank figure slot without creating fake images."""
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.keep_with_next=True
    pf.space_before=Pt(12);pf.space_after=Pt(64)
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(100+number))
    start.set(qn('w:name'),f'figure_pending_{number}');p._p.append(start)
    font(p.add_run(f'图 {number} 待插入'),9)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(100+number));p._p.append(end)


def build():
    # Embed saved figure assets only; never run a model or regenerate plots here.
    QA.mkdir(parents=True,exist_ok=True)
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.left_margin=sec.right_margin=Cm(2.15);sec.top_margin=sec.bottom_margin=Cm(2.05)
    sec.header_distance=sec.footer_distance=Cm(.85)
    for name,size in [('Normal',10.5),('Title',17),('Heading 1',13),('Heading 2',11.5)]:
        st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if name=='Normal' else '黑体')
        st.font.color.rgb=__import__('docx.shared',fromlist=['RGBColor']).RGBColor(0,0,0)
        pf=st.paragraph_format;pf.line_spacing=1.2;pf.space_after=Pt(5);pf.widow_control=True
        pf.keep_with_next=name!='Normal';pf.space_before=Pt(10 if name!='Normal' else 0)
    doc.styles['Normal'].paragraph_format.first_line_indent=Pt(21)
    for border in list(doc.styles._element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    source=SOURCE.read_text(encoding='utf-8');lines=source.splitlines();i=0
    while i<len(lines):
        s=lines[i].strip()
        if not s:i+=1;continue
        if re.fullmatch(r'\[\[FIGURE:\d+\]\]',s):
            add_figure_placeholder(doc,int(re.search(r'\d+',s).group()))
            i+=1;continue
        if s=='$$':
            eq=[];i+=1
            while lines[i].strip()!='$$':eq.append(lines[i].strip());i+=1
            add_equation(doc,' '.join(eq));i+=1;continue
        if s.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{3,}:?',x) for x in cells):rows.append(cells)
                i+=1
            add_table(doc,rows);continue
        if s.startswith('!['):
            alt,path=re.match(r'!\[(.*?)\]\((.*?)\)',s).groups()
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;pf=p.paragraph_format
            pf.first_line_indent=Cm(0);pf.keep_with_next=True
            image_path=SOURCE.parent/path
            from PIL import Image
            with Image.open(image_path) as image:
                pixel_width,pixel_height=image.size
            # Keep a tall panel and its caption on one page without stretching.
            max_height=18.6 if image_path.name == 'pathway_slot_map.png' else 21.2
            width=min(16.5,max_height*pixel_width/pixel_height)
            shape=p.add_run().add_picture(str(image_path),width=Cm(width))
            shape._inline.docPr.set('descr',alt);i+=1;continue
        if s.startswith('#'):
            level=len(s)-len(s.lstrip('#'));p=doc.add_paragraph(style={1:'Title',2:'Heading 1',3:'Heading 2'}[level])
            p.paragraph_format.first_line_indent=Cm(0)
            if level==1:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            title=s[level:].strip()
            if level==1:title=title.replace('与运输感知','\n与运输感知')
            add_inline(p,title,{1:17,2:13,3:11.5}[level],True)
            for r in p.runs:r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'黑体')
        elif s.startswith('*') and s.endswith('*'):
            p=doc.add_paragraph();pf=p.paragraph_format;pf.first_line_indent=Cm(0);pf.keep_together=True;pf.line_spacing=1.1
            add_inline(p,s[1:-1],9)
        else:
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            if re.match(r'^\[\d+\]',s):
                pf=p.paragraph_format;pf.first_line_indent=Cm(-.5);pf.left_indent=Cm(.5);pf.keep_together=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT;add_inline(p,s,9.5)
            else:add_inline(p,s)
        i+=1
    doc.core_properties.title='DCT 面向多模态生存预测的阶段条件最优运输与运输感知通路重建'
    doc.core_properties.subject='v3.13 UNI2-h 论文修订稿 已接入实验图 病例组织图待补'
    doc.core_properties.author=''
    doc.save(OUTPUT)
    pending=[int(x) for x in re.findall(r'\[\[FIGURE:(\d+)\]\]',source)]
    status=json.loads((SOURCE.parent/'V313_MANUSCRIPT_STATUS.json').read_text(encoding='utf-8'))
    assert sorted(pending)==status['figure_placeholders']
    ready=status.get('ready_figures', {})
    assert sorted(pending + [int(k) for k in ready])==list(range(1,8))
    embedded_images=re.findall(r'^!\[.*?\]\((.*?)\)$',source,re.MULTILINE)
    assert len(doc.tables)==8 and len(doc.inline_shapes)==len(embedded_images)
    for entry in list(ready.values())+list(status.get('supplementary_figures',{}).values()):
        for image in entry.get('images',[entry.get('image')]):
            assert image and str(Path(image).relative_to('paper')).replace('\\','/') in embedded_images
    equations=len(list(doc._element.iter(qn('m:oMath'))))
    assert equations==15  # 14 display equations plus one inline stability constant.
    status=json.loads((SOURCE.parent/'V313_MANUSCRIPT_STATUS.json').read_text(encoding='utf-8'))
    report={'docx':str(OUTPUT),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'equations':equations,
            'pending_figures':pending,'partial_figures':status.get('partial_figures',{}),'embedded_images':embedded_images,
            'BLCA':{'mean':BLCA.mean(axis=1).tolist(),'sample_std':BLCA.std(axis=1,ddof=1).tolist()},
            'KIRC':{'mean':KIRC.mean(axis=1).tolist(),'sample_std':KIRC.std(axis=1,ddof=1).tolist()},
            'controls':status['controls'],
            'source_scope':f'70 recorded loss folds; 50 author-supplied Full folds; 20 audited rerun control folds; 10 saved same-model sweep JSONs; 49 aligned DSS literature reference means; OT and UNI2-h sources; no mixed-cohort Overall; {len(pending)} main-figure placeholders; partial tissue panels tracked separately'}
    (QA/'build_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))


if __name__=='__main__':build()
